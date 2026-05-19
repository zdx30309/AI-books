"""Book spreadsheet aggregation for the data analysis tool."""
from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


REPORT_COLUMNS = [
    "时间",
    "星期",
    "AI类",
    "环比",
    "公版类",
    "环比",
    "总时长",
    "环比",
    "汇安AI",
    "汇安公版",
    "悟图AI",
    "悟图公版",
    "精选集/套装书",
]

WEEKDAY_CN = ["星期一", "星期二", "星期三", "星期四", "星期五", "星期六", "星期日"]
COMPANY_ALIASES = {
    "汇安": ("汇安", "惠安"),
    "悟图": ("悟图", "梧图"),
}
COMPANY_COLUMNS = ("公司", "机构", "供应商", "版权方", "出品方", "项目", "归属")
REPORT_HEADER_ALIASES = {
    "时间": "date_key",
    "日期": "date_key",
    "星期": "weekday",
    "AI类时长": "ai_duration",
    "AI类阅读时长": "ai_duration",
    "AI类": "ai_duration",
    "AI 类": "ai_duration",
    "公版类时长": "public_duration",
    "公版类": "public_duration",
    "公版": "public_duration",
    "今日总时长": "total_duration",
    "总阅读时长": "total_duration",
    "总时长": "total_duration",
    "汇安AI": "huian_ai",
    "汇安 AI": "huian_ai",
    "汇安公版": "huian_public",
    "悟图AI": "wutu_ai",
    "悟图 AI": "wutu_ai",
    "悟图公版": "wutu_public",
    "精选集/套装书": "suite_duration",
    "精选集": "suite_duration",
    "套装书": "suite_duration",
    "套装类": "suite_duration",
    "精选集套装书": "suite_duration",
}
PUBLIC_BOOK_RE = re.compile(
    r"古典文学|四大名著|古代诗词|历史典籍|佛教|道教|中医|国学|古籍|经史子集|经典著作"
)


@dataclass
class AnalysisRow:
    date_key: str
    weekday: str
    ai_duration: int | None = None
    ai_mom: float | None = None
    public_duration: int | None = None
    public_mom: float | None = None
    total_duration: int | None = None
    total_mom: float | None = None
    huian_ai: int | None = None
    huian_public: int | None = None
    wutu_ai: int | None = None
    wutu_public: int | None = None
    suite_duration: int | None = None

    def as_dict(self) -> dict[str, Any]:
        return {
            "date_key": self.date_key,
            "weekday": self.weekday,
            "ai_duration": self.ai_duration,
            "ai_mom": self.ai_mom,
            "public_duration": self.public_duration,
            "public_mom": self.public_mom,
            "total_duration": self.total_duration,
            "total_mom": self.total_mom,
            "huian_ai": self.huian_ai,
            "huian_public": self.huian_public,
            "wutu_ai": self.wutu_ai,
            "wutu_public": self.wutu_public,
            "suite_duration": self.suite_duration,
        }


def analyze_book_spreadsheet(filename: str, content: bytes) -> dict[str, Any]:
    return analyze_book_spreadsheets([(filename, content)])


def analyze_book_spreadsheets(files: list[tuple[str, bytes]]) -> dict[str, Any]:
    if not files:
        raise ValueError("请至少上传一个 Excel 文件")

    snapshots = [_analyze_snapshot(filename, content) for filename, content in files]
    dates = [snapshot["report_date"] for snapshot in snapshots]
    report_rows = _blank_month(max(dates)) if _same_month(dates) else _blank_date_range(min(dates), max(dates))
    row_lookup = {row.date_key: row for row in report_rows}
    source_files = []
    total_book_count = 0
    total_matched_count = 0
    detected_fields = set()

    for snapshot in snapshots:
        for metric_date, metrics in snapshot["metrics_by_date"].items():
            key = metric_date.strftime("%Y%m%d")
            target = row_lookup[key]
            _merge_metrics(target, metrics)
        source_files.append(snapshot["source_filename"])
        total_book_count += snapshot["metrics"]["book_count"]
        total_matched_count += snapshot["metrics"]["matched_book_count"]
        if snapshot["metrics"]["company_field_detected"]:
            detected_fields.add(snapshot["metrics"]["company_field_detected"])

    _apply_mom(report_rows)
    summary = {
        "source_filename": "、".join(source_files),
        "source_filenames": source_files,
        "report_date": max(dates).strftime("%Y%m%d"),
        "book_count": total_book_count,
        "matched_book_count": total_matched_count,
        "company_field_detected": "、".join(sorted(detected_fields)) if detected_fields else None,
        "file_count": len(files),
        "notes": _notes(
            {
                "company_field_detected": bool(detected_fields),
                "file_count": len(files),
            }
        ),
    }
    return {
        "summary": summary,
        "rows": [row.as_dict() for row in report_rows],
    }


def _analyze_snapshot(filename: str, content: bytes) -> dict[str, Any]:
    if not content:
        raise ValueError(f"{filename} 上传文件为空")
    if not filename.lower().endswith((".xlsx", ".xlsm")):
        raise ValueError(f"{filename} 不是 .xlsx 或 .xlsm 格式")

    wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"{filename} 中没有可读取的数据")

    headers = [_clean_header(value) for value in rows[0]]
    index = {header: idx for idx, header in enumerate(headers) if header}

    report_metrics = _parse_report_rows(rows, headers, filename)
    if report_metrics:
        return {
            "source_filename": filename,
            "report_date": max(report_metrics),
            "metrics_by_date": report_metrics,
            "metrics": _combine_metrics(report_metrics.values()),
        }

    required = ("今日在读人数", "人均时长（分）")
    missing = [name for name in required if name not in index]
    if missing:
        raise ValueError(f"{filename} 缺少字段：{'、'.join(missing)}")

    report_date = _extract_report_date(filename) or _latest_date(rows[1:], index)
    if not report_date:
        raise ValueError(f"{filename} 无法识别数据日期，请确保文件名包含 8 位日期，例如 20260514")

    current = _aggregate_rows(rows[1:], index)
    return {
        "source_filename": filename,
        "report_date": report_date,
        "metrics_by_date": {report_date: current},
        "metrics": current,
    }


def build_data_analysis_xlsx(result: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "数据分析"
    ws.append(REPORT_COLUMNS)
    for row in result["rows"]:
        ws.append(
            [
                row["date_key"],
                row["weekday"],
                row["ai_duration"],
                row["ai_mom"],
                row["public_duration"],
                row["public_mom"],
                row["total_duration"],
                row["total_mom"],
                row["huian_ai"],
                row["huian_public"],
                row["wutu_ai"],
                row["wutu_public"],
                row["suite_duration"],
            ]
        )

    avg_row = ws.max_row + 1
    ws.cell(avg_row, 1, "日均")
    for col in range(3, 14):
        letter = get_column_letter(col)
        ws.cell(avg_row, col, f'=IFERROR(AVERAGE({letter}2:{letter}{avg_row - 1}),"")')

    _style_workbook(ws, avg_row)
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def safe_data_analysis_filename(source_filename: str) -> str:
    stem = Path(source_filename).stem or "数据分析"
    return f"{stem}_数据分析结果.xlsx"


def _clean_header(value: Any) -> str:
    return str(value or "").strip()


def _extract_report_date(filename: str) -> date | None:
    for match in re.finditer(r"(20\d{6})", filename):
        raw = match.group(1)
        try:
            return datetime.strptime(raw, "%Y%m%d").date()
        except ValueError:
            continue
    return None


def _latest_date(rows: list[tuple[Any, ...]], index: dict[str, int]) -> date | None:
    date_idx = next(
        (index[field] for field in ("上架时间", "日期", "时间") if field in index),
        None,
    )
    if date_idx is None:
        return None
    dates = [_parse_date(row[date_idx]) for row in rows if len(row) > date_idx]
    dates = [value for value in dates if value]
    return max(dates) if dates else None


def _parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in ("%Y.%m.%d", "%Y-%m-%d", "%Y/%m/%d", "%Y%m%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def _same_month(dates: list[date]) -> bool:
    first = dates[0]
    return all(day.year == first.year and day.month == first.month for day in dates)


def _blank_month(report_date: date) -> list[AnalysisRow]:
    first = report_date.replace(day=1)
    next_month = (first.replace(day=28) + timedelta(days=4)).replace(day=1)
    days = (next_month - first).days
    rows = []
    for offset in range(days):
        day = first + timedelta(days=offset)
        rows.append(AnalysisRow(day.strftime("%Y%m%d"), WEEKDAY_CN[day.weekday()]))
    return rows


def _blank_date_range(start_date: date, end_date: date) -> list[AnalysisRow]:
    first = start_date.replace(day=1)
    last_month = end_date.replace(day=1)
    end_next_month = (last_month.replace(day=28) + timedelta(days=4)).replace(day=1)
    days = (end_next_month - first).days
    rows = []
    for offset in range(days):
        day = first + timedelta(days=offset)
        rows.append(AnalysisRow(day.strftime("%Y%m%d"), WEEKDAY_CN[day.weekday()]))
    return rows


def _aggregate_rows(rows: list[tuple[Any, ...]], index: dict[str, int]) -> dict[str, Any]:
    totals: defaultdict[str, int] = defaultdict(int)
    company_field = next((field for field in COMPANY_COLUMNS if field in index), None)
    book_count = 0
    matched_book_count = 0

    for row in rows:
        if not any(value not in (None, "") for value in row):
            continue
        book_count += 1
        readers = _to_number(_value(row, index, "今日在读人数"))
        minutes = _to_number(_value(row, index, "人均时长（分）"))
        duration = int(round(readers * minutes))
        if duration <= 0:
            continue
        matched_book_count += 1
        title = str(_value(row, index, "书籍") or _value(row, index, "书名") or "")
        category = str(_value(row, index, "分类") or "")
        ownership = str(_value(row, index, "出品/非出品") or "")
        company = str(_value(row, index, company_field) or "") if company_field else ""
        bucket = _book_bucket(title, category, ownership)

        totals["total_duration"] += duration
        if bucket == "ai":
            totals["ai_duration"] += duration
        elif bucket == "public":
            totals["public_duration"] += duration
        if _is_suite(title, category):
            totals["suite_duration"] += duration

        company_key = _company_key(company)
        if company_key and bucket in {"ai", "public"}:
            totals[f"{company_key}_{bucket}"] += duration

    return {
        **totals,
        "book_count": book_count,
        "matched_book_count": matched_book_count,
        "company_field_detected": company_field,
    }


def _parse_report_rows(
    rows: list[tuple[Any, ...]],
    headers: list[str],
    filename: str,
) -> dict[date, dict[str, Any]]:
    mapped: dict[str, int] = {}
    company_idx: int | None = None
    for idx, header in enumerate(headers):
        compact = re.sub(r"\s+", "", header)
        field = REPORT_HEADER_ALIASES.get(compact) or REPORT_HEADER_ALIASES.get(header)
        if field:
            mapped[field] = idx
        elif compact in COMPANY_COLUMNS:
            company_idx = idx
        elif any(alias == compact for alias in COMPANY_COLUMNS):
            company_idx = idx

    if "date_key" not in mapped or "total_duration" not in mapped:
        return {}

    # The three 环比 columns sit immediately after AI类、公版类、总时长.
    for field, ring_field in (
        ("ai_duration", "ai_mom"),
        ("public_duration", "public_mom"),
        ("total_duration", "total_mom"),
    ):
        col = mapped.get(field)
        if col is not None and col + 1 < len(headers) and re.sub(r"\s+", "", headers[col + 1]) == "环比":
            mapped[ring_field] = col + 1

    parsed: dict[date, dict[str, Any]] = {}
    for raw in rows[1:]:
        report_date = _parse_date(_cell(raw, mapped.get("date_key")))
        if not report_date:
            continue
        row_metrics = {
            "ai_duration": _int_or_none(_cell(raw, mapped.get("ai_duration"))),
            "ai_mom": _percent_or_none(_cell(raw, mapped.get("ai_mom"))),
            "public_duration": _int_or_none(_cell(raw, mapped.get("public_duration"))),
            "public_mom": _percent_or_none(_cell(raw, mapped.get("public_mom"))),
            "total_duration": _int_or_none(_cell(raw, mapped.get("total_duration"))),
            "total_mom": _percent_or_none(_cell(raw, mapped.get("total_mom"))),
            "huian_ai": _int_or_none(_cell(raw, mapped.get("huian_ai"))),
            "huian_public": _int_or_none(_cell(raw, mapped.get("huian_public"))),
            "wutu_ai": _int_or_none(_cell(raw, mapped.get("wutu_ai"))),
            "wutu_public": _int_or_none(_cell(raw, mapped.get("wutu_public"))),
            "suite_duration": _int_or_none(_cell(raw, mapped.get("suite_duration"))),
            "book_count": 0,
            "matched_book_count": 0,
            "company_field_detected": "日报表",
        }
        if not any(row_metrics.get(key) is not None for key in ("ai_duration", "public_duration", "total_duration")):
            continue

        metrics = parsed.setdefault(
            report_date,
            {
                "book_count": 0,
                "matched_book_count": 0,
                "company_field_detected": "日报表",
            },
        )
        _merge_report_metrics(metrics, row_metrics, str(_cell(raw, company_idx) or ""))
    if parsed:
        for metrics in parsed.values():
            if metrics.get("total_duration") is None:
                total = int(metrics.get("ai_duration") or 0) + int(metrics.get("public_duration") or 0)
                metrics["total_duration"] = total or None
    return parsed


def _merge_report_metrics(target: dict[str, Any], row_metrics: dict[str, Any], company: str) -> None:
    company_key = _company_key(company)
    for field in ("ai_duration", "public_duration", "total_duration", "suite_duration"):
        target[field] = _merge_raw_optional(target.get(field), row_metrics.get(field))
    for field in ("huian_ai", "huian_public", "wutu_ai", "wutu_public"):
        target[field] = _merge_raw_optional(target.get(field), row_metrics.get(field))

    if company_key:
        target[f"{company_key}_ai"] = _merge_raw_optional(
            target.get(f"{company_key}_ai"),
            row_metrics.get("ai_duration"),
        )
        target[f"{company_key}_public"] = _merge_raw_optional(
            target.get(f"{company_key}_public"),
            row_metrics.get("public_duration"),
        )


def _merge_raw_optional(current: Any, next_value: Any) -> int | None:
    if next_value is None:
        return current
    return int(current or 0) + int(next_value or 0)


def _cell(row: tuple[Any, ...], idx: int | None) -> Any:
    if idx is None:
        return None
    return row[idx] if len(row) > idx else None


def _int_or_none(value: Any) -> int | None:
    number = _to_number(value)
    return int(round(number)) if number else None


def _percent_or_none(value: Any) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value or "").strip()
    if not text:
        return None
    number = _to_number(text)
    return number / 100 if "%" in text else number


def _combine_metrics(metrics_list: Any) -> dict[str, Any]:
    combined: defaultdict[str, int] = defaultdict(int)
    company_field_detected = None
    for metrics in metrics_list:
        for field in (
            "ai_duration",
            "public_duration",
            "total_duration",
            "huian_ai",
            "huian_public",
            "wutu_ai",
            "wutu_public",
            "suite_duration",
            "book_count",
            "matched_book_count",
        ):
            combined[field] += int(metrics.get(field) or 0)
        company_field_detected = company_field_detected or metrics.get("company_field_detected")
    return {**combined, "company_field_detected": company_field_detected}


def _value(row: tuple[Any, ...], index: dict[str, int], name: str | None) -> Any:
    if not name or name not in index:
        return None
    idx = index[name]
    return row[idx] if len(row) > idx else None


def _to_number(value: Any) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value or "").strip().replace(",", "")
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    return float(match.group(0)) if match else 0.0


def _book_bucket(title: str, category: str, ownership: str) -> str:
    text = f"{title} {category} {ownership}"
    if "公版" in text or PUBLIC_BOOK_RE.search(text):
        return "public"
    if "出品" in ownership or "ai" in text or "人工智能" in text:
        return "ai"
    return "ai"


def _is_suite(title: str, category: str) -> bool:
    text = f"{title} {category}"
    return bool(re.search(r"精选集|套装|全集|合集|全套|丛书|作品集", text))


def _company_key(company: str) -> str | None:
    for key, aliases in COMPANY_ALIASES.items():
        if any(alias in company for alias in aliases):
            return "huian" if key == "汇安" else "wutu"
    return None


def _apply_metrics(row: AnalysisRow, metrics: dict[str, Any]) -> None:
    row.ai_duration = _none_if_zero(metrics.get("ai_duration"))
    row.public_duration = _none_if_zero(metrics.get("public_duration"))
    row.total_duration = _none_if_zero(metrics.get("total_duration"))
    row.huian_ai = _none_if_zero(metrics.get("huian_ai"))
    row.huian_public = _none_if_zero(metrics.get("huian_public"))
    row.wutu_ai = _none_if_zero(metrics.get("wutu_ai"))
    row.wutu_public = _none_if_zero(metrics.get("wutu_public"))
    row.suite_duration = _none_if_zero(metrics.get("suite_duration"))


def _merge_metrics(row: AnalysisRow, metrics: dict[str, Any]) -> None:
    row.ai_duration = _sum_optional(row.ai_duration, metrics.get("ai_duration"))
    row.public_duration = _sum_optional(row.public_duration, metrics.get("public_duration"))
    row.total_duration = _sum_optional(row.total_duration, metrics.get("total_duration"))
    row.huian_ai = _sum_optional(row.huian_ai, metrics.get("huian_ai"))
    row.huian_public = _sum_optional(row.huian_public, metrics.get("huian_public"))
    row.wutu_ai = _sum_optional(row.wutu_ai, metrics.get("wutu_ai"))
    row.wutu_public = _sum_optional(row.wutu_public, metrics.get("wutu_public"))
    row.suite_duration = _sum_optional(row.suite_duration, metrics.get("suite_duration"))


def _sum_optional(current: int | None, next_value: Any) -> int | None:
    value = int(next_value or 0)
    if value <= 0:
        return current
    return int(current or 0) + value


def _none_if_zero(value: Any) -> int | None:
    value = int(value or 0)
    return value if value else None


def _apply_mom(rows: list[AnalysisRow]) -> None:
    prev_ai = prev_public = prev_total = None
    for row in rows:
        if row.ai_duration is not None and prev_ai:
            row.ai_mom = row.ai_duration / prev_ai - 1
        if row.public_duration is not None and prev_public:
            row.public_mom = row.public_duration / prev_public - 1
        if row.total_duration is not None and prev_total:
            row.total_mom = row.total_duration / prev_total - 1
        prev_ai = row.ai_duration or prev_ai
        prev_public = row.public_duration or prev_public
        prev_total = row.total_duration or prev_total


def _notes(metrics: dict[str, Any]) -> list[str]:
    file_count = int(metrics.get("file_count") or 1)
    notes = ["书籍明细按「今日在读人数 × 人均时长（分）」汇总；日报格式会优先读取已有时长列。"]
    if file_count > 1:
        notes.append(f"已合并 {file_count} 个 Excel 文件；同一天的数据会自动累加。")
    if not metrics.get("company_field_detected"):
        notes.append("未检测到公司/机构字段，汇安与悟图拆分列暂为空。")
    notes.append("AI/公版按公版特征关键词识别，未命中的阅读时长归入 AI 类；套装列按书名/分类中的套装、全集、合集、精选集等关键词判断。")
    return notes


def _style_workbook(ws, avg_row: int) -> None:
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fills = {
        1: "FFFF00",
        2: "FFFF00",
        3: "F2DCDD",
        4: "F2DCDD",
        5: "D9E5F2",
        6: "D9E5F2",
        7: "FFFF00",
        8: "FFFF00",
        9: "9BBB59",
        10: "9BBB59",
        11: "9BBB59",
        12: "9BBB59",
        13: "9BBB59",
    }
    widths = [13, 10, 12, 10, 12, 10, 12, 10, 12, 12, 12, 12, 16]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    for row in ws.iter_rows(min_row=1, max_row=avg_row, max_col=len(REPORT_COLUMNS)):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = Font(name="Songti SC", size=13)
            if cell.row == 1:
                cell.font = Font(name="Songti SC", size=13, bold=True)
                cell.fill = PatternFill("solid", fgColor=header_fills[cell.column])
            elif cell.row == avg_row:
                cell.font = Font(name="Songti SC", size=13, bold=True)
                cell.fill = PatternFill("solid", fgColor="FFFF00" if cell.column <= 8 else "9BBB59")
            elif cell.column in (3, 4):
                cell.fill = PatternFill("solid", fgColor="F2DCDD")
            elif cell.column in (5, 6):
                cell.fill = PatternFill("solid", fgColor="D9E5F2")
            if cell.column in (4, 6, 8):
                cell.number_format = "0.0%"
            elif cell.column >= 3:
                cell.number_format = "0"
    ws.freeze_panes = "A2"
