"""File-backed standalone cover generation jobs."""
from __future__ import annotations

import csv
import json
import re
import shutil
import uuid
import zipfile
from datetime import datetime, timedelta, timezone
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.core.config import PROJECT_ROOT, settings
from app.prompts.goal_driven_templates import get_category_meta, normalize_category

COVER_JOBS_ROOT = PROJECT_ROOT / "backend" / "output" / "cover_jobs"

_INTRO_KEYS = {
    "简介",
    "内容简介",
    "书籍简介",
    "图书简介",
    "intro",
    "summary",
    "description",
    "synopsis",
    "subtitle",
    "brief",
}
_DUPLICATE_STATUSES = {"pending", "processing", "paused", "completed", "failed"}


class CoverDuplicateError(ValueError):
    def __init__(self, duplicates: list[dict[str, Any]], *, window_days: int) -> None:
        self.duplicates = duplicates
        self.window_days = window_days
        names = "、".join(str(item.get("title") or "未命名") for item in duplicates[:6])
        suffix = "等" if len(duplicates) > 6 else ""
        super().__init__(
            f"{window_days} 天内已投喂过相同封面：{names}{suffix}。"
            "请使用已有任务继续生成，或半个月后再投喂。"
        )

_TITLE_KEYS = {"书名", "标题", "图书名", "title", "book_title", "book title", "name"}
_AUTHOR_KEYS = {"作者", "作家", "author", "writer"}
_CATEGORY_KEYS = {"分类", "书架", "题材", "类型", "category", "bookshelf", "type"}


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def create_job_id() -> str:
    return uuid.uuid4().hex


def job_dir(job_id: str) -> Path:
    return COVER_JOBS_ROOT / job_id


def status_path(job_id: str) -> Path:
    return job_dir(job_id) / "status.json"


def rows_path(job_id: str) -> Path:
    return job_dir(job_id) / "rows.json"


def _clean_cell(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _normalized_key(value: str) -> str:
    return value.strip().lower().replace(" ", "_")


def _row_value(row: dict[str, str], keys: set[str]) -> str:
    lowered = {_normalized_key(k): v for k, v in row.items()}
    for key in keys:
        value = lowered.get(_normalized_key(key))
        if value:
            return value.strip()
    return ""


def row_title(row: dict[str, str]) -> str:
    return _row_value(row, _TITLE_KEYS)


def row_author(row: dict[str, str]) -> str:
    return _row_value(row, _AUTHOR_KEYS)


def row_category(row: dict[str, str]) -> str:
    return _row_value(row, _CATEGORY_KEYS)


def row_intro(row: dict[str, str]) -> str:
    return _row_value(row, _INTRO_KEYS)


def normalize_cover_category(category: str | None, fallback: str | None = None) -> tuple[str, str]:
    raw = (category or fallback or "social_science").strip() or "social_science"
    try:
        normalized = normalize_category(raw)
    except ValueError:
        normalized = normalize_category(fallback or "social_science")
    meta = get_category_meta(normalized)
    return normalized, str(meta["label"])


def _with_common_keys(row: dict[str, str], default_category: str | None = None) -> dict[str, str]:
    next_row = dict(row)
    title = row_title(row)
    author = row_author(row)
    intro = row_intro(row)
    category, category_label = normalize_cover_category(row_category(row), default_category)
    if title:
        next_row.setdefault("title", title)
    if author:
        next_row.setdefault("author", author)
    if intro:
        next_row.setdefault("intro", intro)
    next_row["category"] = category
    next_row["category_label"] = category_label
    return next_row


def _parse_csv(content: bytes, default_category: str | None = None) -> list[dict[str, str]]:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(StringIO(text))
    rows: list[dict[str, str]] = []
    for raw in reader:
        row = {str(k).strip(): _clean_cell(v) for k, v in raw.items() if k and str(k).strip()}
        if any(row.values()):
            rows.append(_with_common_keys(row, default_category))
    return rows


def _parse_xlsx(content: bytes, default_category: str | None = None) -> list[dict[str, str]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    try:
        headers = [_clean_cell(cell) for cell in next(iterator)]
    except StopIteration:
        return []
    headers = [header or f"column_{idx + 1}" for idx, header in enumerate(headers)]
    rows: list[dict[str, str]] = []
    for values in iterator:
        row = {
            headers[idx]: _clean_cell(values[idx] if idx < len(values) else "")
            for idx in range(len(headers))
            if headers[idx]
        }
        if any(row.values()):
            rows.append(_with_common_keys(row, default_category))
    workbook.close()
    return rows


def parse_cover_spreadsheet(
    filename: str,
    content: bytes,
    *,
    default_category: str | None = None,
) -> list[dict[str, str]]:
    normalized_default, _ = normalize_cover_category(default_category)
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        rows = _parse_csv(content, normalized_default)
    elif suffix in {".xlsx", ".xlsm"}:
        rows = _parse_xlsx(content, normalized_default)
    else:
        raise ValueError("请上传 .xlsx、.xlsm 或 .csv 文件")

    valid = [row for row in rows if row_title(row)]
    if not valid:
        raise ValueError("Excel 至少需要一列书名，列名可为「书名」或 title")
    max_rows = max(1, int(settings.COVER_JOB_MAX_ROWS or 500))
    if len(valid) > max_rows:
        raise ValueError(f"一次最多生成 {max_rows} 张封面")
    return valid


def _string_row(row: dict[str, Any]) -> dict[str, str]:
    return {str(k): str(v) for k, v in row.items() if v is not None}


def _fingerprint_text(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "").strip().lower())


def cover_fingerprint(row: dict[str, Any]) -> str:
    normalized_row = _string_row(row)
    title = row_title(normalized_row)
    author = row_author(normalized_row)
    return "|".join((_fingerprint_text(title), _fingerprint_text(author)))


def _parse_iso_datetime(value: Any) -> datetime | None:
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except ValueError:
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(timezone.utc)


def _duplicate_summary(row: dict[str, Any], *, source: str, job_id: str | None = None) -> dict[str, Any]:
    normalized_row = _string_row(row)
    return {
        "title": row_title(normalized_row) or str(row.get("title") or ""),
        "author": row_author(normalized_row) or str(row.get("author") or ""),
        "category": str(row.get("category") or row_category(normalized_row) or ""),
        "category_label": str(row.get("category_label") or ""),
        "source": source,
        "job_id": job_id,
    }


def find_recent_duplicate_covers(
    owner_id: str,
    rows: list[dict[str, str]],
    *,
    window_days: int | None = None,
) -> list[dict[str, Any]]:
    days = max(0, int(window_days if window_days is not None else settings.COVER_DUPLICATE_WINDOW_DAYS or 0))
    if days <= 0:
        return []

    duplicates: list[dict[str, Any]] = []
    seen: dict[str, dict[str, str]] = {}
    incoming: dict[str, dict[str, str]] = {}
    for row in rows:
        key = cover_fingerprint(row)
        if not key.strip("|"):
            continue
        if key in seen:
            duplicates.append(_duplicate_summary(row, source="current_upload"))
            continue
        seen[key] = row
        incoming[key] = row

    if not COVER_JOBS_ROOT.exists() or not incoming:
        return duplicates

    cutoff = datetime.now(timezone.utc) - timedelta(days=days)
    for path in COVER_JOBS_ROOT.glob("*/status.json"):
        try:
            status = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            continue
        if status.get("owner_id") != owner_id:
            continue
        if str(status.get("status") or "") not in _DUPLICATE_STATUSES:
            continue
        created_at = _parse_iso_datetime(status.get("created_at") or status.get("updated_at"))
        if created_at is not None and created_at < cutoff:
            continue

        job_id = str(status.get("job_id") or path.parent.name)
        for item in status.get("items") or []:
            if not isinstance(item, dict):
                continue
            metadata = item.get("metadata") if isinstance(item.get("metadata"), dict) else item
            key = cover_fingerprint(metadata)
            if key in incoming:
                duplicates.append(_duplicate_summary(incoming[key], source="recent_job", job_id=job_id))

    deduped: dict[tuple[str, str | None], dict[str, Any]] = {}
    for item in duplicates:
        deduped[(str(item.get("title") or ""), item.get("job_id"))] = item
    return list(deduped.values())


def safe_image_name(index: int, title: str) -> str:
    safe = re.sub(r'[\\/:*?"<>|\r\n]+', "_", title).strip(" ._") or f"cover_{index}"
    return f"{index:03d}_{safe[:60]}.png"


def read_rows(job_id: str) -> list[dict[str, str]]:
    path = rows_path(job_id)
    if not path.exists():
        raise FileNotFoundError(job_id)
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, list):
        raise ValueError(f"cover job rows are invalid: {job_id}")
    return [dict(row) for row in data if isinstance(row, dict)]


def read_status(job_id: str) -> dict[str, Any]:
    path = status_path(job_id)
    if not path.exists():
        raise FileNotFoundError(job_id)
    return json.loads(path.read_text(encoding="utf-8"))


def write_status(job_id: str, data: dict[str, Any]) -> None:
    directory = job_dir(job_id)
    directory.mkdir(parents=True, exist_ok=True)
    data["updated_at"] = now_iso()
    tmp = status_path(job_id).with_suffix(".json.tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(status_path(job_id))


def pause_cover_job(job_id: str, status: dict[str, Any] | None = None) -> dict[str, Any]:
    data = dict(status or read_status(job_id))
    if data.get("status") in {"completed", "deleted"}:
        return data
    for item in data.get("items") or []:
        if item.get("status") == "processing":
            item["status"] = "pending"
            item["error"] = None
    data["status"] = "paused"
    data["current_step"] = "已暂停"
    write_status(job_id, data)
    return data


def resume_cover_job(job_id: str, status: dict[str, Any] | None = None) -> dict[str, Any]:
    data = dict(status or read_status(job_id))
    if data.get("status") == "completed":
        return data
    if data.get("status") == "deleted":
        raise FileNotFoundError(job_id)
    data["status"] = "pending"
    data["current_step"] = "等待继续生成"
    write_status(job_id, data)
    return data


def delete_cover_job(job_id: str) -> None:
    directory = job_dir(job_id)
    if not directory.exists():
        raise FileNotFoundError(job_id)
    shutil.rmtree(directory)


def list_cover_jobs(owner_id: str, *, limit: int = 50) -> list[dict[str, Any]]:
    if not COVER_JOBS_ROOT.exists():
        return []
    jobs: list[dict[str, Any]] = []
    for path in COVER_JOBS_ROOT.glob("*/status.json"):
        try:
            status = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            continue
        if status.get("owner_id") == owner_id:
            jobs.append(status)
    jobs.sort(key=lambda item: str(item.get("updated_at") or item.get("created_at") or ""), reverse=True)
    return jobs[: max(1, int(limit or 50))]


def create_cover_job(
    owner_id: str,
    filename: str,
    rows: list[dict[str, str]],
    *,
    default_category: str | None = None,
) -> dict[str, Any]:
    duplicates = find_recent_duplicate_covers(
        owner_id,
        rows,
        window_days=settings.COVER_DUPLICATE_WINDOW_DAYS,
    )
    if duplicates:
        raise CoverDuplicateError(duplicates, window_days=settings.COVER_DUPLICATE_WINDOW_DAYS)

    job_id = create_job_id()
    directory = job_dir(job_id)
    directory.mkdir(parents=True, exist_ok=True)
    rows_path(job_id).write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    created_at = now_iso()
    default_category_id, default_category_label = normalize_cover_category(default_category)
    items = [
        {
            "index": idx,
            "title": row_title(row),
            "author": row_author(row) or None,
            "category": str(row.get("category") or default_category_id),
            "category_label": str(row.get("category_label") or default_category_label),
            "status": "pending",
            "image_url": None,
            "error": None,
            "metadata": row,
        }
        for idx, row in enumerate(rows, start=1)
    ]
    data: dict[str, Any] = {
        "job_id": job_id,
        "owner_id": owner_id,
        "status": "pending",
        "progress": 0,
        "total": len(rows),
        "completed": 0,
        "failed": 0,
        "current_step": "排队中",
        "source_filename": filename,
        "default_category": default_category_id,
        "default_category_label": default_category_label,
        "created_at": created_at,
        "updated_at": created_at,
        "items": items,
    }
    write_status(job_id, data)
    return data


def image_path(job_id: str, index: int, title: str) -> Path:
    return job_dir(job_id) / "images" / safe_image_name(index, title)


def build_zip(job_id: str, status: dict[str, Any]) -> bytes:
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        written = 0
        for item in status.get("items") or []:
            if item.get("status") != "completed":
                continue
            path = image_path(job_id, int(item["index"]), str(item.get("title") or "cover"))
            if path.exists() and path.stat().st_size > 0:
                archive.write(path, arcname=path.name)
                written += 1
        if written == 0:
            raise FileNotFoundError("没有可下载的封面")
    return buffer.getvalue()


__all__ = [
    "build_zip",
    "CoverDuplicateError",
    "cover_fingerprint",
    "create_cover_job",
    "delete_cover_job",
    "find_recent_duplicate_covers",
    "image_path",
    "list_cover_jobs",
    "normalize_cover_category",
    "parse_cover_spreadsheet",
    "read_rows",
    "read_status",
    "pause_cover_job",
    "resume_cover_job",
    "row_author",
    "row_category",
    "row_intro",
    "row_title",
    "write_status",
]
