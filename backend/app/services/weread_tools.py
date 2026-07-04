"""Utilities for WeRead recommendation exports."""
from __future__ import annotations

import csv
import html
import json
import os
import re
import zipfile
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, quote, urljoin, urlparse
from xml.sax.saxutils import escape

import httpx
from openpyxl import load_workbook


WEREAD_HOST = "https://weread.qq.com"
WEREAD_I_HOST = "https://i.weread.qq.com"
REPO_ROOT = Path(__file__).resolve().parents[3]
PROJECT_INDEX_PATH = Path(__file__).resolve().parent / "reference_assets" / "weread_project_index.json"
WEREAD_COMPANY_INDEX_PATH = Path(
    os.getenv("WEREAD_COMPANY_INDEX_PATH", str(REPO_ROOT / ".local" / "weread_company_index.json"))
)
USER_AGENT = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"
)
_PROJECT_INDEX_CACHE: list[dict[str, Any]] | None = None
_COMPANY_INDEX_CACHE: dict[str, Any] | None = None
_EXCEL_TITLE_HEADERS = ("书名", "原始书名", "标题", "原始标题", "图书", "作品", "项目名称")
_EXCEL_AUTHOR_HEADERS = ("原著作者", "原始作者", "作者", "匹配作者")
_EXCEL_URL_HEADERS = ("微信读书链接", "微信链接", "来源链接", "原始链接", "匹配链接", "链接", "url")
_EXCEL_NEW_TITLE_HEADERS = ("新书名/上架后改名", "新书名", "上架后改名")
_EXCEL_MATCHED_TITLE_HEADERS = ("匹配书名", "匹配标题", "匹配图书")
_NON_BOOK_TITLE_RE = re.compile(r"(?:表格|书单|项目进度|进度汇总|数据汇总|汇总表)")


@dataclass
class CatalogBook:
    title: str
    author: str | None = None
    source_url: str | None = None


@dataclass
class CompanyLookupInput:
    company_name: str | None
    seed_items: list[CatalogBook]


@dataclass
class WeReadBookResult:
    query_title: str
    query_author: str | None
    matched_title: str | None = None
    matched_author: str | None = None
    recommendation: float | None = None
    rating_count: str | None = None
    today_reading: str | None = None
    total_reading: str | None = None
    word_count: str | None = None
    reading_count: str | None = None
    comment_count: str | None = None
    company_name: str | None = None
    copyright_id: str | None = None
    translator: str | None = None
    source_url: str | None = None
    profit: float | None = None
    project_editor: str | None = None
    project_title: str | None = None
    project_word_count_10k: str | None = None
    project_finished_at: str | None = None
    project_summary: str | None = None
    project_pen_name: str | None = None
    project_new_title: str | None = None
    project_original_author: str | None = None
    project_recommendation: float | None = None
    project_comment_count: str | None = None
    project_tags: str | None = None
    project_reading_count: str | None = None
    project_weread_url: str | None = None
    project_negative_review: str | None = None
    project_parent_record: str | None = None
    project_status: str | None = None
    project_progress: str | None = None
    project_update_date: str | None = None
    project_start_date: str | None = None
    project_time_span: str | None = None
    project_note: str | None = None
    project_source_sheet: str | None = None
    project_match_status: str | None = None
    status: str = "pending"
    error: str | None = None

    def as_dict(self) -> dict[str, Any]:
        return {
            "query_title": self.query_title,
            "query_author": self.query_author,
            "matched_title": self.matched_title,
            "matched_author": self.matched_author,
            "recommendation": self.recommendation,
            "rating_count": self.rating_count,
            "today_reading": self.today_reading,
            "total_reading": self.total_reading,
            "word_count": self.word_count,
            "reading_count": self.reading_count,
            "comment_count": self.comment_count,
            "company_name": self.company_name,
            "copyright_id": self.copyright_id,
            "translator": self.translator,
            "source_url": self.source_url,
            "profit": self.profit,
            "project_editor": self.project_editor,
            "project_title": self.project_title,
            "project_word_count_10k": self.project_word_count_10k,
            "project_finished_at": self.project_finished_at,
            "project_summary": self.project_summary,
            "project_pen_name": self.project_pen_name,
            "project_new_title": self.project_new_title,
            "project_original_author": self.project_original_author,
            "project_recommendation": self.project_recommendation,
            "project_comment_count": self.project_comment_count,
            "project_tags": self.project_tags,
            "project_reading_count": self.project_reading_count,
            "project_weread_url": self.project_weread_url,
            "project_negative_review": self.project_negative_review,
            "project_parent_record": self.project_parent_record,
            "project_status": self.project_status,
            "project_progress": self.project_progress,
            "project_update_date": self.project_update_date,
            "project_start_date": self.project_start_date,
            "project_time_span": self.project_time_span,
            "project_note": self.project_note,
            "project_source_sheet": self.project_source_sheet,
            "project_match_status": self.project_match_status,
            "status": self.status,
            "error": self.error,
        }


def normalize_title(value: str | None) -> str:
    text = (value or "").strip().lower()
    text = re.sub(r"[《》〈〉“”\"'（）()\[\]【】\s·•\-_:：，,。.!！?？]", "", text)
    return text


def parse_catalog(text: str) -> list[CatalogBook]:
    items: list[CatalogBook] = []
    seen: set[str] = set()
    header: list[str] | None = None
    known_editors = _known_project_editors()

    for raw_line in (text or "").splitlines():
        line = raw_line.strip().strip("|")
        if not line:
            continue
        row = _split_row(line)
        if _looks_like_catalog_header(row):
            header = row
            continue

        url_match = re.search(r"https?://[^\s,，]+", line)
        source_url = url_match.group(0) if url_match else None
        clean_line = re.sub(r"^\s*\d+[\.\)、)]\s*", "", line)
        clean_line = re.sub(r"https?://[^\s,，]+", "", clean_line).strip(" ,，")

        title = ""
        author: str | None = None
        if header and row:
            title = _value_from_header(row, header, (*_EXCEL_TITLE_HEADERS, "title"))
            if not title:
                title = _value_from_header(row, header, (*_EXCEL_NEW_TITLE_HEADERS, *_EXCEL_MATCHED_TITLE_HEADERS))
            author = _value_from_header(row, header, (*_EXCEL_AUTHOR_HEADERS, "author")) or None
            source_url = source_url or _value_from_header(row, header, _EXCEL_URL_HEADERS) or None

        bracket_match = re.search(r"《([^》]+)》", clean_line)
        if not title and bracket_match:
            title = bracket_match.group(1).strip()
            rest = clean_line[bracket_match.end() :].strip(" -—,，\t")
            author = rest or None
        elif not title:
            row = _split_row(clean_line)
            if row:
                title, author = _title_author_from_row(row, known_editors)
            elif source_url:
                title = source_url

        title = _clean_cell(title)
        author = _clean_cell(author) if author else None
        if normalize_title(title) in known_editors and not source_url:
            continue
        if not title and source_url:
            title = source_url
        if not title:
            continue

        key = normalize_title(title) or source_url or title
        if key in seen:
            continue
        seen.add(key)
        items.append(CatalogBook(title=title, author=author, source_url=source_url))

    return items


def parse_catalog_workbook(content: bytes) -> list[CatalogBook]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("无法读取 Excel 文件，请确认上传的是 .xlsx 表格") from exc

    items: list[CatalogBook] = []
    seen: set[str] = set()
    seen_titles: set[str] = set()
    for sheet in workbook.worksheets:
        header: list[str] | None = None
        header_row = 0
        for row_index, raw_row in enumerate(sheet.iter_rows(values_only=True), start=1):
            row = [_excel_cell_text(value) for value in raw_row]
            if not any(row):
                continue
            if header is None:
                if _looks_like_excel_catalog_header(row):
                    header = row
                    header_row = row_index
                continue
            if row_index <= header_row:
                continue

            title = _value_from_header(row, header, _EXCEL_TITLE_HEADERS)
            if not title:
                title = _value_from_header(row, header, _EXCEL_NEW_TITLE_HEADERS)
            if not title:
                title = _value_from_header(row, header, _EXCEL_MATCHED_TITLE_HEADERS)
            author = _value_from_header(row, header, _EXCEL_AUTHOR_HEADERS) or None
            source_url = _value_from_header(row, header, _EXCEL_URL_HEADERS) or None
            source_url = _extract_first_url(source_url)
            if not title and source_url:
                title = source_url
            title = _clean_cell(title)
            author = _clean_cell(author) if author else None
            if not title:
                continue
            if _looks_like_non_book_excel_row(title, author, source_url):
                continue

            key = _catalog_seen_key(title, author, source_url)
            title_key = normalize_title(title)
            if key in seen or (title_key and title_key in seen_titles):
                continue
            seen.add(key)
            if title_key:
                seen_titles.add(title_key)
            items.append(CatalogBook(title=title, author=author, source_url=source_url))

    if not items:
        raise ValueError("Excel 中没有识别到可抓取的书名或微信读书链接")
    return items


def catalog_books_to_text(items: list[CatalogBook]) -> str:
    lines = ["书名\t原著作者\t微信读书链接"]
    for item in items:
        lines.append("\t".join((item.title, item.author or "", item.source_url or "")))
    return "\n".join(lines)


def _excel_cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    if text.endswith(".0") and re.fullmatch(r"\d+\.0", text):
        return text[:-2]
    return text


def _looks_like_excel_catalog_header(row: list[str]) -> bool:
    normalized = [normalize_title(cell) for cell in row if cell]
    has_title = any(
        normalize_title(name) in normalized
        for name in (*_EXCEL_TITLE_HEADERS, *_EXCEL_NEW_TITLE_HEADERS, *_EXCEL_MATCHED_TITLE_HEADERS)
    )
    has_context = any(
        normalize_title(name) in normalized
        for name in (
            *_EXCEL_URL_HEADERS,
            *_EXCEL_AUTHOR_HEADERS,
            "编辑",
            "编辑名",
            "索引字段（导出时备份用，勿动）",
            "简介",
            "项目状态",
            "推荐值",
            "版权方（公司名字）",
            "版权方",
            "阅读人数",
            "评论数量",
            "阅读",
        )
    )
    return has_title and has_context


def _extract_first_url(value: str | None) -> str | None:
    match = re.search(r"https?://[^\s，,。；;)）]+", value or "")
    return match.group(0) if match else None


def _catalog_seen_key(title: str, author: str | None, source_url: str | None) -> str:
    url_key = _project_url_key(source_url)
    if url_key:
        return url_key
    return f"{normalize_title(title)}|{normalize_title(author)}"


def _looks_like_non_book_excel_row(title: str, author: str | None, source_url: str | None) -> bool:
    normalized = normalize_title(title)
    if source_url:
        return False
    if not normalized:
        return True
    if normalized in _known_project_editors():
        return True
    if _NON_BOOK_TITLE_RE.search(title):
        return True
    return normalized in {
        "书名",
        "原始书名",
        "匹配书名",
        "标题",
        "项目状态",
        "当前大概进度",
        "索引字段导出时备份用勿动",
    }


def parse_profit_table(text: str) -> dict[str, float]:
    profits: dict[str, float] = {}
    for raw_line in (text or "").splitlines():
        line = raw_line.strip().strip("|")
        if not line:
            continue
        if re.search(r"^(书名|标题|title).*(利润|profit)", line, re.I):
            continue

        row = _split_row(line)
        title = ""
        profit: float | None = None
        if len(row) >= 2:
            title = row[0]
            profit = _parse_float(row[-1])

        if profit is None:
            match = re.search(r"(.+?)[\s,，:：]+([+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?)\s*(?:元|¥|￥)?$", line)
            if match:
                title = match.group(1).strip()
                profit = _parse_float(match.group(2))

        if title and profit is not None:
            profits[normalize_title(title)] = profit
    return profits


async def analyze_weread_books(
    catalog_text: str,
    profit_text: str | None = None,
    limit: int | None = None,
) -> list[WeReadBookResult]:
    catalog = parse_catalog(catalog_text)
    return await analyze_weread_catalog(catalog, profit_text=profit_text, limit=limit)


async def analyze_weread_workbook(
    content: bytes,
    profit_text: str | None = None,
    limit: int | None = None,
) -> list[WeReadBookResult]:
    catalog = parse_catalog_workbook(content)
    return await analyze_weread_catalog(catalog, profit_text=profit_text, limit=limit)


async def analyze_weread_catalog(
    catalog: list[CatalogBook],
    profit_text: str | None = None,
    limit: int | None = None,
) -> list[WeReadBookResult]:
    if not catalog:
        return []

    profits = parse_profit_table(profit_text or "")
    timeout = httpx.Timeout(12.0, connect=8.0)
    headers = {"User-Agent": USER_AGENT, "Accept-Language": "zh-CN,zh;q=0.9"}
    async with httpx.AsyncClient(timeout=timeout, headers=headers, follow_redirects=True) as client:
        results: list[WeReadBookResult] = []
        selected_catalog = catalog if limit is None else catalog[: max(0, limit)]
        for item in selected_catalog:
            result = WeReadBookResult(query_title=item.title, query_author=item.author)
            try:
                if item.source_url:
                    result = await fetch_reader_result(client, item)
                else:
                    result = await fetch_search_result(client, item)
            except Exception as exc:  # pragma: no cover - defensive network boundary
                result.status = "failed"
                result.error = str(exc)[:240]

            result.profit = _match_profit(profits, result)
            _attach_project_match(result)
            _upsert_company_index_from_result(result)
            results.append(result)
        return results


async def analyze_weread_company_books(company_name: str, limit: int | None = None) -> list[WeReadBookResult]:
    lookup = _parse_company_lookup_input(company_name)
    target = _valid_company_name(lookup.company_name)
    max_books = max(1, limit or 500)
    timeout = httpx.Timeout(16.0, connect=8.0)
    headers = {"User-Agent": USER_AGENT, "Accept-Language": "zh-CN,zh;q=0.9"}
    async with httpx.AsyncClient(timeout=timeout, headers=headers, follow_redirects=True) as client:
        seed_details: list[WeReadBookResult] = []
        if lookup.seed_items:
            for seed in lookup.seed_items[:10]:
                if not seed.source_url:
                    continue
                try:
                    detail = await fetch_reader_result(client, seed)
                except Exception:
                    continue
                if detail.company_name or detail.copyright_id:
                    seed_details.append(detail)
                    _upsert_company_index_from_result(detail)
                    if not target and detail.company_name:
                        target = detail.company_name

        if not target:
            return []

        target_ids = _target_company_ids_from_index(target)
        target_ids.update(
            str(detail.copyright_id)
            for detail in seed_details
            if detail.copyright_id and _company_matches(detail.company_name, target)
        )
        if target.isdigit():
            target_ids.add(target)

        candidates = await _collect_company_book_candidates(
            client,
            target,
            max_books=max_books,
            seed_items=lookup.seed_items,
            target_ids=target_ids,
        )
        results: list[WeReadBookResult] = []
        processed: set[str] = set()
        queued: set[str] = {
            key for key in (_candidate_seen_key(card) for card in candidates) if key
        }
        accepted: set[str] = set()
        expanded_keywords: set[str] = set()
        target_ids.update(
            str(card.get("copyright_id"))
            for card in candidates
            if (
                _company_matches(card.get("company_name"), target)
                or card.get("source_kind") == "seed_search"
            )
            and card.get("copyright_id")
        )

        index = 0
        while index < len(candidates) and len(results) < max_books:
            card = candidates[index]
            index += 1
            if len(results) >= max_books:
                break
            key = _candidate_seen_key(card)
            if not key or key in processed:
                continue
            processed.add(key)

            result = _result_from_company_card(card, target)
            detail: WeReadBookResult | None = None
            if card.get("source_url"):
                try:
                    detail = await fetch_reader_result(
                        client,
                        CatalogBook(
                            title=card.get("title") or target,
                            author=card.get("author"),
                            source_url=card.get("source_url"),
                        ),
                    )
                except Exception as exc:  # pragma: no cover - public network boundary
                    result.status = "partial"
                    result.error = f"详情页补充失败：{str(exc)[:120]}"

            if detail:
                result = _merge_company_detail(result, detail)
                if result.copyright_id and _company_matches(result.company_name, target):
                    target_ids.add(result.copyright_id)

            if _company_result_matches(result, target, target_ids):
                _attach_project_match(result)
                result.company_name = result.company_name or target
                result.comment_count = _comment_count_for_display(result.comment_count or result.rating_count)
                result.status = "matched" if result.recommendation is not None else "partial"
                if result.status == "matched":
                    result.error = None
                _upsert_company_index_from_result(result)
                result_key = _company_result_key(result)
                if result_key and result_key in accepted:
                    continue
                if result_key:
                    accepted.add(result_key)
                results.append(result)

                for keyword in _company_expansion_keywords(result):
                    if keyword in expanded_keywords or len(candidates) >= max_books * 4:
                        continue
                    expanded_keywords.add(keyword)
                    await _add_company_cards(
                        candidates,
                        queued,
                        await _fetch_search_book_cards(client, keyword, max_books=40),
                    )

                if result.copyright_id:
                    await _add_company_cards(
                        candidates,
                        queued,
                        await _fetch_private_copyright_book_cards(
                            client,
                            result.copyright_id,
                            max_books=max_books,
                        ),
                    )

        return results


async def fetch_search_result(client: httpx.AsyncClient, item: CatalogBook) -> WeReadBookResult:
    url = f"{WEREAD_HOST}/web/search/books?keyword={quote(item.title)}"
    response = await client.get(url)
    response.raise_for_status()
    html_text = response.text
    cards = _parse_search_cards(html_text)
    if not cards:
        return WeReadBookResult(
            query_title=item.title,
            query_author=item.author,
            source_url=url,
            status="not_found",
            error="微信读书搜索页没有返回匹配书籍",
        )

    selected = _select_card(cards, item)
    if not selected:
        return WeReadBookResult(
            query_title=item.title,
            query_author=item.author,
            source_url=url,
            status="not_found",
            error="搜索结果中没有作者对得上的书籍",
        )
    result = WeReadBookResult(
        query_title=item.title,
        query_author=item.author,
        matched_title=selected.get("title"),
        matched_author=selected.get("author"),
        recommendation=_parse_float(selected.get("recommendation")),
        today_reading=selected.get("today_reading"),
        source_url=selected.get("source_url") or url,
        status="matched",
    )
    if selected.get("source_url"):
        try:
            detail = await fetch_reader_result(
                client,
                CatalogBook(
                    title=item.title,
                    author=item.author,
                    source_url=selected.get("source_url"),
                ),
            )
            result.matched_title = detail.matched_title or result.matched_title
            result.matched_author = detail.matched_author or result.matched_author
            if not _author_matches(item.author, result.matched_author):
                return WeReadBookResult(
                    query_title=item.title,
                    query_author=item.author,
                    matched_title=result.matched_title,
                    matched_author=result.matched_author,
                    source_url=result.source_url,
                    status="not_found",
                    error="匹配作者不一致，已跳过该搜索结果",
                )
            result.recommendation = detail.recommendation if detail.recommendation is not None else result.recommendation
            result.rating_count = detail.rating_count or result.rating_count
            result.total_reading = detail.total_reading or result.total_reading
            result.word_count = detail.word_count or result.word_count
            result.company_name = detail.company_name or result.company_name
            result.translator = detail.translator or result.translator
            result.source_url = detail.source_url or result.source_url
            result.status = "matched" if result.recommendation is not None else detail.status
            result.error = detail.error if result.status != "matched" else None
        except Exception:
            pass
    return result


async def fetch_reader_result(client: httpx.AsyncClient, item: CatalogBook) -> WeReadBookResult:
    response = await client.get(item.source_url or "")
    response.raise_for_status()
    body = response.text
    title = _first_match(body, [
        r'"name"\s*:\s*"([^"]+)"',
        r'<meta property="og:title" content="([^"]+)"',
        r"<title>(.*?)\s*-\s*.*?微信读书</title>",
    ]) or _extract_nuxt_book_field(body, "title")
    author = _first_match(body, [
        r'"author"\s*:\s*\{[^}]*"name"\s*:\s*"([^"]+)"',
        r"<title>.*?-\s*(.*?)\s*-\s*微信读书</title>",
    ]) or _extract_nuxt_book_field(body, "author")
    recommendation = _extract_recommendation_percent(body)
    rating_count = _extract_rating_count(body)
    total_reading = _extract_stat(body, "阅读")
    word_count = _extract_stat(body, "字数")
    company_name = _extract_copyright_holder(body)
    copyright_id = _extract_copyright_id(body)
    translator = _extract_translator(body)

    clean_title = _clean_text(title) or item.title
    clean_author = _clean_text(author) or item.author
    author_mismatch = item.source_url is None and not _author_matches(item.author, clean_author)
    return WeReadBookResult(
        query_title=item.title,
        query_author=item.author,
        matched_title=clean_title,
        matched_author=clean_author,
        recommendation=recommendation,
        rating_count=rating_count or "0人点评",
        total_reading=total_reading,
        word_count=word_count,
        company_name=company_name,
        copyright_id=copyright_id,
        translator=translator,
        source_url=item.source_url,
        status="not_found" if author_mismatch else ("matched" if recommendation is not None else "partial"),
        error=(
            "匹配作者不一致，已跳过该搜索结果"
            if author_mismatch
            else (None if recommendation is not None else "已读取书籍页，但未找到推荐值")
        ),
    )


async def _collect_company_book_candidates(
    client: httpx.AsyncClient,
    company_name: str,
    *,
    max_books: int,
    seed_items: list[CatalogBook] | None = None,
    target_ids: set[str] | None = None,
) -> list[dict[str, Any]]:
    candidates: list[dict[str, Any]] = []
    seen: set[str] = set()

    async def add_cards(cards: list[dict[str, Any]]) -> None:
        await _add_company_cards(candidates, seen, cards)

    await add_cards(_company_seed_cards(seed_items or []))
    for seed in seed_items or []:
        if seed.source_url:
            continue
        seed_title = _clean_text(seed.title)
        if not seed_title:
            continue
        seed_cards = await _fetch_search_book_cards(client, seed_title, max_books=min(20, max_books))
        for card in seed_cards:
            card["source_kind"] = "seed_search"
        await add_cards(seed_cards)
    await add_cards(_company_candidate_cards_from_cache(company_name, target_ids=target_ids or set(), max_books=max_books))
    await add_cards(await _fetch_publisher_book_cards(client, company_name, max_books=max_books))
    await add_cards(_company_candidate_cards_from_project_index(company_name, max_books=max_books))

    if len(candidates) < max_books:
        for keyword in _company_search_keywords(company_name):
            await add_cards(await _fetch_search_book_cards(client, keyword, max_books=min(40, max_books)))
            if len(candidates) >= max_books:
                break

    return candidates[:max_books]


async def _add_company_cards(
    candidates: list[dict[str, Any]],
    seen: set[str],
    cards: list[dict[str, Any]],
) -> None:
    for card in cards:
        key = _candidate_seen_key(card)
        if not key or key in seen:
            continue
        seen.add(key)
        candidates.append(card)


def _company_seed_cards(seed_items: list[CatalogBook]) -> list[dict[str, Any]]:
    cards: list[dict[str, Any]] = []
    for item in seed_items:
        if not item.source_url:
            continue
        cards.append(
            {
                "title": item.title or item.source_url,
                "author": item.author,
                "source_url": item.source_url,
                "source_kind": "seed",
            }
        )
    return cards


async def _fetch_publisher_book_cards(
    client: httpx.AsyncClient,
    keyword: str,
    *,
    max_books: int,
) -> list[dict[str, Any]]:
    cards: list[dict[str, Any]] = []
    max_idx = 0
    page_size = min(40, max(1, max_books))
    for _ in range(20):
        url = (
            f"{WEREAD_HOST}/web/search/publisher_books"
            f"?keyword={quote(keyword)}&maxIdx={max_idx}&count={page_size}"
        )
        response = await client.get(url)
        response.raise_for_status()
        payload = response.json()
        page_cards = _parse_api_book_cards(payload)
        if not page_cards:
            break
        for card in page_cards:
            card.setdefault("company_name", keyword)
            card.setdefault("source_kind", "publisher")
        cards.extend(page_cards)
        if len(cards) >= max_books or not payload.get("hasMore"):
            break
        next_idx = max(_parse_int(card.get("search_idx")) or 0 for card in page_cards)
        if next_idx <= max_idx:
            break
        max_idx = next_idx
    return cards[:max_books]


async def _fetch_search_book_cards(
    client: httpx.AsyncClient,
    keyword: str,
    *,
    max_books: int,
) -> list[dict[str, Any]]:
    url = f"{WEREAD_HOST}/web/search/global?keyword={quote(keyword)}"
    response = await client.get(url)
    response.raise_for_status()
    try:
        cards = _parse_api_book_cards(response.json())
    except json.JSONDecodeError:
        cards = _parse_search_cards(response.text)
    for card in cards:
        card.setdefault("source_kind", "search")
    return cards[:max_books]


async def _fetch_private_copyright_book_cards(
    client: httpx.AsyncClient,
    copyright_id: str,
    *,
    max_books: int,
) -> list[dict[str, Any]]:
    cookie = _weread_auth_cookie()
    if not cookie or not copyright_id:
        return []

    cards: list[dict[str, Any]] = []
    max_idx = 0
    page_size = min(50, max(1, max_books))
    headers = {
        "Cookie": cookie,
        "User-Agent": USER_AGENT,
        "Accept": "application/json,text/plain,*/*",
        "Referer": WEREAD_HOST,
    }
    for _ in range(30):
        urls = (
            f"{WEREAD_I_HOST}/book/copyright?cpid={quote(copyright_id)}&maxIdx={max_idx}&count={page_size}",
            f"{WEREAD_I_HOST}/book/copyright?copyrightId={quote(copyright_id)}&maxIdx={max_idx}&count={page_size}",
        )
        page_cards: list[dict[str, Any]] = []
        payload: Any = None
        for url in urls:
            response = await client.get(url, headers=headers)
            if response.status_code in {401, 403, 404}:
                continue
            response.raise_for_status()
            try:
                payload = response.json()
            except json.JSONDecodeError:
                payload = {}
            page_cards = _parse_api_book_cards(payload)
            if page_cards:
                break
        if not page_cards:
            break
        for card in page_cards:
            card.setdefault("copyright_id", copyright_id)
            card.setdefault("source_kind", "private_copyright")
        cards.extend(page_cards)
        if len(cards) >= max_books or not _payload_has_more(payload):
            break
        next_idx = max(_parse_int(card.get("search_idx")) or 0 for card in page_cards)
        if next_idx <= max_idx:
            max_idx += len(page_cards)
        else:
            max_idx = next_idx
    return cards[:max_books]


def _company_candidate_cards_from_cache(
    company_name: str,
    *,
    target_ids: set[str],
    max_books: int,
) -> list[dict[str, Any]]:
    index = _load_company_index()
    company_keys = {normalize_title(company_name)}
    company_keys.update(
        normalize_title(index.get("copyright_ids", {}).get(str(copyright_id)))
        for copyright_id in target_ids
        if index.get("copyright_ids", {}).get(str(copyright_id))
    )
    cards: list[dict[str, Any]] = []
    for company_key in company_keys:
        if not company_key:
            continue
        company = index.get("companies", {}).get(company_key)
        if not isinstance(company, dict):
            continue
        for raw_card in (company.get("books") or {}).values():
            if not isinstance(raw_card, dict):
                continue
            card = dict(raw_card)
            card.setdefault("company_name", company.get("company_name") or company_name)
            card.setdefault("source_kind", "company_cache")
            cards.append(card)
            if len(cards) >= max_books:
                return cards
    return cards


def _company_candidate_cards_from_project_index(company_name: str, *, max_books: int) -> list[dict[str, Any]]:
    cards: list[dict[str, Any]] = []
    for record in _load_project_records():
        if not _company_matches(record.get("company_name"), company_name):
            continue
        url = _project_text(record, "project_weread_url")
        title = _project_text(record, "title")
        if not title and not url:
            continue
        cards.append(
            {
                "title": title or url,
                "author": _project_text(record, "original_author") or _project_text(record, "pen_name"),
                "source_url": url,
                "company_name": _project_text(record, "company_name"),
                "recommendation": _project_recommendation_value(record.get("project_recommendation")),
                "rating_count": _normalize_comment_count(record.get("project_comment_count")),
                "reading_count": _project_text(record, "project_reading_count"),
                "source_kind": "project_index",
            }
        )
        if len(cards) >= max_books:
            break
    return cards


def _parse_company_lookup_input(raw_value: str) -> CompanyLookupInput:
    target: str | None = None
    seed_items: list[CatalogBook] = []
    for raw_line in str(raw_value or "").splitlines():
        line = _clean_text(raw_line)
        if not line:
            continue
        url = _extract_first_url(line)
        if url:
            title = _clean_text(line.replace(url, " ").strip(" ，,。；;"))
            if title and not target:
                target = title
            seed_items.append(CatalogBook(title=title or url, source_url=url))
            continue
        if not target:
            target = line
        else:
            seed_items.append(CatalogBook(title=line))
    return CompanyLookupInput(company_name=target, seed_items=seed_items)


def _target_company_ids_from_index(company_name: str) -> set[str]:
    index = _load_company_index()
    company = index.get("companies", {}).get(normalize_title(company_name))
    if not isinstance(company, dict):
        return set()
    return {
        str(value)
        for value in company.get("copyright_ids") or []
        if str(value).strip()
    }


def _company_expansion_keywords(result: WeReadBookResult) -> list[str]:
    title = _clean_text(result.matched_title or result.query_title)
    if not title:
        return []
    stems = [
        title,
        re.split(r"[（(：:：·•\s]", title, maxsplit=1)[0],
        re.split(r"[第上下中前后续]", title, maxsplit=1)[0],
    ]
    if len(title) >= 6:
        stems.append(title[:6])
    if len(title) >= 4:
        stems.append(title[:4])
    keywords: list[str] = []
    for stem in stems:
        clean = _clean_text(stem).strip("《》“”\"'：:，,。；;、")
        if len(clean) < 2 or clean in keywords:
            continue
        keywords.append(clean)
    return keywords[:5]


def _company_search_keywords(company_name: str) -> list[str]:
    text = _clean_text(company_name)
    keywords = [text]
    if text.endswith(("文化", "出版", "传媒", "互众", "互动", "工作室")) and len(text) > 4:
        keywords.append(text[:-2])
    return [keyword for index, keyword in enumerate(keywords) if keyword and keyword not in keywords[:index]]


def _parse_api_book_cards(payload: Any) -> list[dict[str, Any]]:
    if not isinstance(payload, dict):
        return []
    raw_books = _api_book_rows(payload)
    cards: list[dict[str, Any]] = []
    for raw in raw_books:
        if not isinstance(raw, dict):
            continue
        book = raw.get("bookInfo") if isinstance(raw.get("bookInfo"), dict) else raw
        title = _clean_text(str(book.get("title") or ""))
        if not title:
            continue
        author = _clean_text(str(book.get("author") or "")) or None
        source_url = _book_source_url(book)
        copyright_info = book.get("copyrightInfo") if isinstance(book.get("copyrightInfo"), dict) else {}
        company_name = _valid_company_name(copyright_info.get("name") if copyright_info else None)
        publisher = _valid_company_name(book.get("publisher"))
        copyright_id = (
            str(copyright_info.get("id"))
            if copyright_info and copyright_info.get("id") not in (None, "")
            else str(book.get("cpid")) if book.get("cpid") not in (None, "", 0) else None
        )
        cards.append(
            {
                "title": title,
                "author": author,
                "source_url": source_url,
                "company_name": company_name or publisher,
                "copyright_id": copyright_id,
                "recommendation": _rating_to_percent(book.get("newRating")),
                "rating_count": _normalize_comment_count(book.get("newRatingCount")),
                "today_reading": _format_people_count(raw.get("readingCount") or raw.get("reading")),
                "reading_count": _format_people_count(
                    raw.get("totalReadingCount")
                    or raw.get("readCount")
                    or raw.get("readingTotal")
                    or raw.get("readingCount")
                    or raw.get("reading")
                ),
                "search_idx": raw.get("searchIdx"),
            }
        )
    return cards


def _api_book_rows(payload: dict[str, Any]) -> list[Any]:
    for key in ("books", "bookInfos", "items", "list"):
        value = payload.get(key)
        if isinstance(value, list):
            return value
    data = payload.get("data")
    if isinstance(data, dict):
        for key in ("books", "bookInfos", "items", "list"):
            value = data.get(key)
            if isinstance(value, list):
                return value
    return []


def _payload_has_more(payload: Any) -> bool:
    if not isinstance(payload, dict):
        return False
    for key in ("hasMore", "has_more", "more", "hasNext"):
        value = payload.get(key)
        if isinstance(value, bool):
            return value
        if isinstance(value, (int, str)):
            return str(value).lower() in {"1", "true", "yes"}
    data = payload.get("data")
    if isinstance(data, dict):
        return _payload_has_more(data)
    return False


def _book_source_url(book: dict[str, Any]) -> str | None:
    deep_link = str(book.get("deepLink") or "")
    if deep_link.startswith("http"):
        return deep_link.replace("\\u002F", "/")
    encode_id = str(book.get("encodeId") or "")
    if encode_id:
        path = "bookDetail" if book.get("soldout") or book.get("autoBookRequest") else "reader"
        return f"{WEREAD_HOST}/web/{path}/{encode_id}"
    book_id = str(book.get("bookId") or "")
    if book_id:
        return f"{WEREAD_HOST}/web/bookDetail/{book_id}"
    return None


def _result_from_company_card(card: dict[str, Any], company_name: str) -> WeReadBookResult:
    title = _clean_text(str(card.get("title") or company_name))
    return WeReadBookResult(
        query_title=title,
        query_author=card.get("author"),
        matched_title=title,
        matched_author=card.get("author"),
        recommendation=_parse_float(card.get("recommendation")),
        rating_count=card.get("rating_count"),
        today_reading=card.get("today_reading"),
        total_reading=card.get("reading_count"),
        reading_count=card.get("reading_count"),
        word_count=card.get("word_count"),
        comment_count=_comment_count_for_display(card.get("rating_count")),
        company_name=card.get("company_name"),
        copyright_id=str(card.get("copyright_id")) if card.get("copyright_id") else None,
        translator=card.get("translator"),
        source_url=card.get("source_url"),
        status="partial",
    )


def _merge_company_detail(base: WeReadBookResult, detail: WeReadBookResult) -> WeReadBookResult:
    base.matched_title = detail.matched_title or base.matched_title
    base.matched_author = detail.matched_author or base.matched_author
    base.recommendation = detail.recommendation if detail.recommendation is not None else base.recommendation
    base.rating_count = detail.rating_count or base.rating_count
    base.total_reading = detail.total_reading or base.total_reading
    base.reading_count = detail.reading_count or detail.total_reading or base.reading_count
    base.word_count = detail.word_count or base.word_count
    base.company_name = detail.company_name or base.company_name
    base.copyright_id = detail.copyright_id or base.copyright_id
    base.translator = detail.translator or base.translator
    base.source_url = detail.source_url or base.source_url
    base.comment_count = _comment_count_for_display(detail.comment_count or detail.rating_count or base.comment_count)
    return base


def _company_result_matches(result: WeReadBookResult, company_name: str, target_ids: set[str]) -> bool:
    if _company_matches(result.company_name, company_name):
        return True
    return bool(result.copyright_id and result.copyright_id in target_ids)


def _company_matches(value: Any, target: str) -> bool:
    wanted = normalize_title(target)
    actual = normalize_title(str(value or ""))
    if not wanted or not actual:
        return False
    return wanted == actual or wanted in actual or actual in wanted


def _candidate_seen_key(card: dict[str, Any]) -> str:
    url_key = _project_url_key(card.get("source_url"))
    if url_key:
        return url_key
    return "|".join(
        part
        for part in (
            normalize_title(card.get("title")),
            normalize_title(card.get("author")),
        )
        if part
    )


def _company_result_key(result: WeReadBookResult) -> str:
    return "|".join(
        part
        for part in (
            normalize_title(result.matched_title or result.query_title),
            normalize_title(result.matched_author or result.query_author),
            str(result.copyright_id or ""),
        )
        if part
    )


def _rating_to_percent(value: Any) -> float | None:
    number = _parse_float(value)
    if number is None:
        return None
    return number / 10 if number > 100 else number


def _format_people_count(value: Any) -> str | None:
    number = _parse_float(value)
    if number is None:
        return None
    if number >= 10000:
        return f"{number / 10000:.1f}万人".replace(".0万", "万")
    return f"{int(number)}人"


def _parse_int(value: Any) -> int | None:
    try:
        return int(str(value))
    except (TypeError, ValueError):
        return None


def build_weread_xlsx(rows: list[dict[str, Any]]) -> bytes:
    headers = [
        "原始书名",
        "原始链接",
        "匹配书名",
        "匹配作者",
        "译者",
        "版权方（公司名字）",
        "推荐值",
        "阅读人数",
        "评论数量",
    ]
    body_rows = [headers]
    for row in rows:
        body_rows.append(
            [
                row.get("query_title"),
                _source_link_for_export(row),
                row.get("matched_title"),
                row.get("matched_author"),
                row.get("translator"),
                _company_for_display(row.get("company_name")),
                _format_percent(row.get("recommendation")),
                row.get("reading_count"),
                _comment_count_for_display(row.get("comment_count")),
            ]
        )

    shared_strings: list[str] = []
    shared_index: dict[str, int] = {}

    def sst_index(value: Any) -> int:
        text = "" if value is None else str(value)
        if text not in shared_index:
            shared_index[text] = len(shared_strings)
            shared_strings.append(text)
        return shared_index[text]

    sheet_rows: list[str] = []
    for r_idx, row in enumerate(body_rows, start=1):
        cells: list[str] = []
        for c_idx, value in enumerate(row, start=1):
            ref = f"{_col_name(c_idx)}{r_idx}"
            if isinstance(value, (int, float)) and value is not None:
                cells.append(f'<c r="{ref}"><v>{value}</v></c>')
            else:
                cells.append(f'<c r="{ref}" t="s"><v>{sst_index(value)}</v></c>')
        sheet_rows.append(f'<row r="{r_idx}">{"".join(cells)}</row>')

    shared_xml = "".join(
        f"<si><t>{escape(value)}</t></si>" for value in shared_strings
    )
    sheet_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <cols>
    <col min="1" max="4" width="24" customWidth="1"/>
    <col min="5" max="9" width="18" customWidth="1"/>
  </cols>
  <sheetData>{"".join(sheet_rows)}</sheetData>
</worksheet>"""

    output = BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", _content_types_xml())
        zf.writestr("_rels/.rels", _root_rels_xml())
        zf.writestr("xl/workbook.xml", _workbook_xml())
        zf.writestr("xl/_rels/workbook.xml.rels", _workbook_rels_xml())
        zf.writestr("xl/worksheets/sheet1.xml", sheet_xml)
        zf.writestr(
            "xl/sharedStrings.xml",
            f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="{len(shared_strings)}" uniqueCount="{len(shared_strings)}">{shared_xml}</sst>""",
        )
        zf.writestr("xl/styles.xml", _styles_xml())
        zf.writestr("docProps/core.xml", _core_xml())
        zf.writestr("docProps/app.xml", _app_xml())
    return output.getvalue()


def _source_link_for_export(row: dict[str, Any]) -> str | None:
    for key in ("source_url", "project_weread_url", "query_author"):
        value = row.get(key)
        url = _extract_first_url(str(value)) if value is not None else None
        if url:
            return url
    return None


def _company_for_display(value: Any) -> str:
    return _valid_company_name(str(value or "")) or "无内容"


def _comment_count_for_display(value: Any) -> str:
    return _normalize_comment_count(value) or "0人点评"


def _weread_auth_cookie() -> str:
    return (
        os.getenv("WEREAD_COOKIE")
        or os.getenv("WEREAD_AUTH_COOKIE")
        or ""
    ).strip()


def _parse_search_cards(html_text: str) -> list[dict[str, str | None]]:
    cards: list[dict[str, str | None]] = []
    for match in re.finditer(r'<li class="wr_bookList_item">([\s\S]*?)</li>', html_text):
        block = match.group(1)
        href = _first_match(block, [r'<a href="([^"]+)" class="wr_bookList_item_link"'])
        title = _first_match(block, [r'<p class="wr_bookList_item_title">([\s\S]*?)</p>'])
        author = _first_match(block, [r'<p class="wr_bookList_item_author">([\s\S]*?)</p>'])
        recommendation = _first_match(block, [r'wr_bookList_item_reading_percent">([^<]+)</span>'])
        today = _first_match(block, [r'wr_bookList_item_reading_number">([^<]+)</span>人今日阅读'])
        source_url = urljoin(WEREAD_HOST, href) if href else None
        cards.append(
            {
                "title": _clean_text(title),
                "author": _clean_text(author),
                "recommendation": _clean_text(recommendation),
                "today_reading": _clean_text(today),
                "source_url": source_url,
            }
        )
    return cards


def _select_card(cards: list[dict[str, str | None]], item: CatalogBook) -> dict[str, str | None] | None:
    wanted_title = normalize_title(item.title)
    wanted_author = normalize_title(item.author) if _query_author_requires_match(item.author) else ""
    for card in cards:
        if normalize_title(card.get("title")) == wanted_title:
            if not wanted_author or _author_matches(item.author, card.get("author")):
                return card
    for card in cards:
        card_title = normalize_title(card.get("title"))
        if wanted_title and (wanted_title in card_title or card_title in wanted_title):
            if wanted_author and not _author_matches(item.author, card.get("author")):
                continue
            return card
    if wanted_author:
        return None
    return cards[0] if cards else None


def _query_author_requires_match(value: str | None) -> bool:
    text = _clean_cell(value)
    return bool(text and not _extract_first_url(text))


def _author_matches(query_author: str | None, matched_author: str | None) -> bool:
    if not _query_author_requires_match(query_author):
        return True
    wanted = normalize_title(query_author)
    actual = normalize_title(matched_author)
    if not wanted or not actual:
        return False
    return wanted == actual or wanted in actual or actual in wanted


def _match_profit(profits: dict[str, float], result: WeReadBookResult) -> float | None:
    keys = [
        normalize_title(result.query_title),
        normalize_title(result.matched_title),
    ]
    for key in keys:
        if key and key in profits:
            return profits[key]
    return None


def _attach_project_match(result: WeReadBookResult) -> None:
    record, match_status = _find_project_record(result)
    if record:
        result.company_name = result.company_name or _project_text(record, "company_name")
        result.project_editor = _project_text(record, "editor")
        result.project_title = _project_text(record, "title")
        result.project_word_count_10k = _project_text(record, "word_count_10k")
        result.project_finished_at = _project_text(record, "finished_at")
        result.project_summary = _project_text(record, "summary")
        result.project_pen_name = _project_text(record, "pen_name")
        result.project_new_title = _project_text(record, "new_title")
        result.project_original_author = _project_text(record, "original_author")
        result.project_recommendation = _project_recommendation_value(record.get("project_recommendation"))
        result.project_comment_count = _project_text(record, "project_comment_count")
        result.project_tags = _project_text(record, "project_tags")
        result.project_reading_count = _project_text(record, "project_reading_count")
        result.project_weread_url = _project_text(record, "project_weread_url")
        result.project_negative_review = _project_text(record, "negative_review")
        result.project_parent_record = _project_text(record, "parent_record")
        result.project_status = _project_text(record, "project_status")
        result.project_progress = _project_text(record, "project_progress")
        result.project_update_date = _project_text(record, "project_update_date")
        result.project_start_date = _project_text(record, "project_start_date")
        result.project_time_span = _project_text(record, "time_span")
        result.project_note = _project_text(record, "project_note")
        result.project_source_sheet = _project_text(record, "source_sheet")
        result.project_match_status = match_status
        if result.recommendation is None:
            result.recommendation = result.project_recommendation
        if not result.rating_count:
            result.rating_count = _normalize_comment_count(result.project_comment_count)
        if not result.total_reading:
            result.total_reading = result.project_reading_count
        if not result.source_url:
            result.source_url = result.project_weread_url
        if result.status in {"pending", "failed", "not_found"}:
            result.status = "matched" if result.recommendation is not None else "partial"
            result.error = None

    result.reading_count = result.total_reading or result.project_reading_count
    result.comment_count = _comment_count_for_display(result.rating_count or result.project_comment_count)


def _find_project_record(result: WeReadBookResult) -> tuple[dict[str, Any] | None, str | None]:
    records = _load_project_records()
    url_keys = {
        key
        for key in (
            _project_url_key(result.source_url),
            _project_url_key(result.query_title),
            _project_url_key(result.matched_title),
        )
        if key
    }
    if url_keys:
        for record in records:
            if record.get("_url_key") in url_keys:
                return record, "链接匹配"

    title_keys = {
        key
        for key in (
            normalize_title(result.query_title),
            normalize_title(result.matched_title),
        )
        if key
    }
    if title_keys:
        for record in records:
            if record.get("_title_key") in title_keys:
                return record, "书名匹配"
    return None, None


def _load_project_records() -> list[dict[str, Any]]:
    global _PROJECT_INDEX_CACHE
    if _PROJECT_INDEX_CACHE is not None:
        return _PROJECT_INDEX_CACHE
    if not PROJECT_INDEX_PATH.exists():
        _PROJECT_INDEX_CACHE = []
        return _PROJECT_INDEX_CACHE
    try:
        payload = json.loads(PROJECT_INDEX_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        _PROJECT_INDEX_CACHE = []
        return _PROJECT_INDEX_CACHE

    records: list[dict[str, Any]] = []
    for raw in payload.get("records") or []:
        if not isinstance(raw, dict):
            continue
        record = dict(raw)
        record["_url_key"] = _project_url_key(record.get("project_weread_url"))
        record["_title_key"] = normalize_title(record.get("title"))
        records.append(record)
    _PROJECT_INDEX_CACHE = records
    return records


def _load_company_index() -> dict[str, Any]:
    global _COMPANY_INDEX_CACHE
    if _COMPANY_INDEX_CACHE is not None:
        return _COMPANY_INDEX_CACHE
    if not WEREAD_COMPANY_INDEX_PATH.exists():
        _COMPANY_INDEX_CACHE = {"version": 1, "companies": {}, "copyright_ids": {}}
        return _COMPANY_INDEX_CACHE
    try:
        payload = json.loads(WEREAD_COMPANY_INDEX_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        payload = {}
    if not isinstance(payload, dict):
        payload = {}
    payload.setdefault("version", 1)
    payload.setdefault("companies", {})
    payload.setdefault("copyright_ids", {})
    _COMPANY_INDEX_CACHE = payload
    return _COMPANY_INDEX_CACHE


def _save_company_index(index: dict[str, Any]) -> None:
    try:
        WEREAD_COMPANY_INDEX_PATH.parent.mkdir(parents=True, exist_ok=True)
        tmp_path = WEREAD_COMPANY_INDEX_PATH.with_suffix(".tmp")
        tmp_path.write_text(
            json.dumps(index, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        tmp_path.replace(WEREAD_COMPANY_INDEX_PATH)
    except OSError:
        return


def _upsert_company_index_from_result(result: WeReadBookResult) -> None:
    company_name = _valid_company_name(result.company_name)
    if not company_name:
        return
    title = _clean_text(result.matched_title or result.query_title)
    source_url = result.source_url or result.project_weread_url
    if not title and not source_url:
        return

    index = _load_company_index()
    company_key = normalize_title(company_name)
    companies = index.setdefault("companies", {})
    company = companies.setdefault(
        company_key,
        {
            "company_name": company_name,
            "copyright_ids": [],
            "books": {},
        },
    )
    company["company_name"] = company_name
    if result.copyright_id:
        copyright_id = str(result.copyright_id)
        copyright_ids = list(company.get("copyright_ids") or [])
        if copyright_id not in copyright_ids:
            copyright_ids.append(copyright_id)
        company["copyright_ids"] = copyright_ids
        index.setdefault("copyright_ids", {})[copyright_id] = company_name

    card = {
        "title": title,
        "author": _clean_text(result.matched_author or result.query_author) or None,
        "source_url": source_url,
        "company_name": company_name,
        "copyright_id": str(result.copyright_id) if result.copyright_id else None,
        "recommendation": result.recommendation,
        "rating_count": _comment_count_for_display(result.comment_count or result.rating_count),
        "reading_count": result.reading_count or result.total_reading,
        "word_count": result.word_count,
        "translator": result.translator,
        "updated_at": datetime.now(UTC).isoformat(timespec="seconds").replace("+00:00", "Z"),
    }
    book_key = _company_cache_book_key(card)
    if not book_key:
        return
    books = company.setdefault("books", {})
    existing = books.get(book_key) if isinstance(books.get(book_key), dict) else {}
    books[book_key] = {**existing, **{key: value for key, value in card.items() if value not in (None, "")}}
    index["updated_at"] = datetime.now(UTC).isoformat(timespec="seconds").replace("+00:00", "Z")
    _save_company_index(index)


def _company_cache_book_key(card: dict[str, Any]) -> str:
    url_key = _project_url_key(card.get("source_url"))
    if url_key:
        return url_key
    return "|".join(
        part
        for part in (
            normalize_title(card.get("title")),
            normalize_title(card.get("author")),
        )
        if part
    )


def _known_project_editors() -> set[str]:
    return {
        normalize_title(record.get("editor"))
        for record in _load_project_records()
        if normalize_title(record.get("editor"))
    }


def _project_url_key(value: Any) -> str:
    text = str(value or "").strip().strip(" ，,。；;)")
    if not text or "http" not in text:
        return ""
    match = re.search(r"https?://[^\s，,。；;）)]+", text)
    if match:
        text = match.group(0)
    reader = re.search(r"/web/reader/([^/?#]+)", text)
    if reader:
        return f"reader:{reader.group(1)}"
    parsed = urlparse(text)
    if not parsed.netloc:
        return ""
    if parsed.path.rstrip("/") == "/book-detail":
        query = parse_qs(parsed.query)
        detail_id = next(
            (
                value.strip()
                for key in ("v", "bookId", "bookid")
                for value in (query.get(key) or [])
                if value and value.strip()
            ),
            "",
        )
        if detail_id:
            return f"book-detail:{detail_id}"
    return f"{parsed.netloc.lower()}{parsed.path.rstrip('/')}?{parsed.query}".rstrip("?")


def _project_text(record: dict[str, Any], key: str) -> str | None:
    value = record.get(key)
    text = str(value).strip() if value is not None else ""
    return text or None


def _project_recommendation_value(value: Any) -> float | None:
    number = _parse_float(value)
    if number is None:
        return None
    if 0 < number <= 1:
        return number * 100
    return number


def _split_row(line: str) -> list[str]:
    if not line:
        return []
    normalized = line.replace("，", ",")
    if "\t" in normalized:
        row = normalized.split("\t")
    elif "," in normalized:
        try:
            row = next(csv.reader(StringIO(normalized)))
        except csv.Error:
            row = re.split(r"[\t,]", normalized)
    else:
        row = re.split(r"\s{2,}|[|｜]", normalized)
    return [_clean_cell(item) for item in row if _clean_cell(item)]


def _looks_like_catalog_header(row: list[str]) -> bool:
    if not row:
        return False
    normalized = [normalize_title(cell) for cell in row]
    title_headers = {normalize_title(name) for name in (*_EXCEL_TITLE_HEADERS, *_EXCEL_NEW_TITLE_HEADERS, *_EXCEL_MATCHED_TITLE_HEADERS, "title")}
    context_headers = {
        normalize_title(name)
        for name in (
            *_EXCEL_AUTHOR_HEADERS,
            *_EXCEL_URL_HEADERS,
            "author",
            "编辑",
            "编辑名",
            "版权方（公司名字）",
            "推荐值",
            "阅读人数",
            "评论数量",
        )
    }
    return any(cell in title_headers for cell in normalized) and any(
        cell in context_headers
        for cell in normalized
    )


def _value_from_header(row: list[str], header: list[str], names: tuple[str, ...]) -> str:
    wanted = {normalize_title(name) for name in names}
    for index, name in enumerate(header):
        if normalize_title(name) in wanted and index < len(row):
            return row[index]
    return ""


def _title_author_from_row(row: list[str], known_editors: set[str]) -> tuple[str, str | None]:
    if not row:
        return "", None
    first = row[0]
    if len(row) >= 2 and normalize_title(first) in known_editors:
        return row[1], row[2] if len(row) > 2 else None
    return first, row[1] if len(row) > 1 else None


def _clean_cell(value: str | None) -> str:
    return (value or "").strip().strip('"').strip("'").strip()


def _clean_text(value: str | None) -> str:
    if not value:
        return ""
    text = re.sub(r"<[^>]+>", "", value)
    return html.unescape(text).replace("\xa0", " ").strip()


def _parse_float(value: Any) -> float | None:
    if value is None:
        return None
    text = str(value).strip().replace("%", "").replace(",", "").replace("¥", "").replace("￥", "")
    try:
        return float(text)
    except ValueError:
        return None


def _first_match(text: str, patterns: list[str]) -> str | None:
    for pattern in patterns:
        match = re.search(pattern, text, re.S)
        if match:
            return html.unescape(match.group(1)).strip()
    return None


def _extract_nuxt_book_field(text: str, field: str) -> str | None:
    start = text.find("bookInfo:{")
    if start < 0:
        return None
    block = text[start : start + 8000]
    match = re.search(rf"\b{re.escape(field)}\s*:\s*([^,\}}]+)", block)
    if not match:
        return None

    token = match.group(1).strip()
    literal = _decode_js_string_literal(token)
    if literal:
        return literal

    values = _nuxt_argument_values(text)
    value = values.get(token)
    return value or None


def _nuxt_argument_values(text: str) -> dict[str, str]:
    script_start = text.find("window.__NUXT__=(function(")
    if script_start < 0:
        return {}
    script_end = text.find("</script>", script_start)
    script = text[script_start : script_end if script_end >= 0 else len(text)]

    params_match = re.search(r"window\.__NUXT__=\(function\(([^)]*)\)", script)
    args_start = script.rfind("}(")
    args_end = script.rfind("));")
    if not params_match or args_start < 0 or args_end < args_start:
        return {}

    params = [part.strip() for part in params_match.group(1).split(",") if part.strip()]
    args = _split_js_arguments(script[args_start + 2 : args_end])
    values: dict[str, str] = {}
    for param, raw_value in zip(params, args):
        value = _decode_js_string_literal(raw_value.strip())
        if value:
            values[param] = value
    return values


def _split_js_arguments(text: str) -> list[str]:
    args: list[str] = []
    current: list[str] = []
    quote_char = ""
    escaped = False
    depth = 0
    for char in text:
        current.append(char)
        if quote_char:
            if escaped:
                escaped = False
            elif char == "\\":
                escaped = True
            elif char == quote_char:
                quote_char = ""
            continue
        if char in {"'", '"'}:
            quote_char = char
        elif char in "([{":
            depth += 1
        elif char in ")]}":
            depth = max(0, depth - 1)
        elif char == "," and depth == 0:
            args.append("".join(current[:-1]).strip())
            current = []
    tail = "".join(current).strip()
    if tail:
        args.append(tail)
    return args


def _decode_js_string_literal(value: str) -> str | None:
    text = value.strip()
    if len(text) < 2 or text[0] not in {"'", '"'} or text[-1] != text[0]:
        return None
    if text[0] == "'":
        text = '"' + text[1:-1].replace('"', r"\"") + '"'
    try:
        decoded = json.loads(text)
    except json.JSONDecodeError:
        return _clean_text(value.strip("'\""))
    return _clean_text(decoded)


def _extract_copyright_holder(text: str) -> str | None:
    priority_values = _json_string_values(
        text,
        (
            "copyrightInfo.name",
            "copyrightHolder",
            "copyrightOwner",
            "copyrightCompany",
            "copyrightName",
            "copyright",
            "rightsHolder",
            "rightsOwner",
            "rightsName",
            "ownerName",
            "cpName",
            "companyName",
        ),
    )
    for value in priority_values:
        clean_value = _valid_company_name(value)
        if clean_value:
            return clean_value

    for value in _json_string_values(text, ("publisher", "publisherName", "press")):
        clean_value = _valid_company_name(value)
        if clean_value:
            return clean_value
    return None


def _extract_copyright_id(text: str) -> str | None:
    for pattern in (
        r'"copyrightInfo"\s*:\s*\{[^{}]*"id"\s*:\s*([0-9]+)',
        r"\bcopyrightInfo\s*:\s*\{[^{}]*\bid\s*:\s*([0-9]+)",
        r'"cpid"\s*:\s*([0-9]+)',
        r"\bcpid\s*:\s*([0-9]+)",
    ):
        match = re.search(pattern, text)
        if match and match.group(1) != "0":
            return match.group(1)
    return None


def _extract_translator(text: str) -> str | None:
    for value in _json_string_values(text, ("translator",)):
        if value:
            return value

    seg_block = _first_match(text, [r'"translatorSeg"\s*:\s*\[([\s\S]*?)\]'])
    seg_values = [
        _clean_text(_decode_json_string(value))
        for value in re.findall(r'"words"\s*:\s*"([^"]*)"', seg_block or "")
    ]
    joined = "".join(value for value in seg_values if value)
    return joined or None


def _json_string_values(text: str, keys: tuple[str, ...]) -> list[str]:
    values: list[str] = []
    for key in keys:
        if key == "copyrightInfo.name":
            patterns = (
                r'"copyrightInfo"\s*:\s*\{[^{}]*"name"\s*:\s*"([^"]*)"',
                r"\bcopyrightInfo\s*:\s*\{[^{}]*\bname\s*:\s*\"([^\"]*)\"",
            )
        else:
            escaped = re.escape(key)
            patterns = (
                rf'"{escaped}"\s*:\s*"([^"]*)"',
                rf"\b{escaped}\s*:\s*\"([^\"]*)\"",
            )
        for pattern in patterns:
            for raw_value in re.findall(pattern, text, re.S):
                value = _clean_text(_decode_json_string(raw_value))
                if value:
                    values.append(value)
    return values


def _valid_company_name(value: str | None) -> str | None:
    text = _clean_text(value).strip(" ，,。；;:：")
    if not text:
        return None
    if text.lower() in {"null", "none", "undefined", "true", "false"}:
        return None
    if len(text) == 1 and re.fullmatch(r"[a-zA-Z]", text):
        return None
    return text


def _extract_recommendation_percent(text: str) -> float | None:
    direct = _first_match(
        text,
        [
            r"微信读书推荐值\s*([0-9.]+%)",
            r"book_rating_item_label_number[\s\S]{0,160}?([0-9]+(?:\.[0-9]+)?)\s*<span[^>]*>%</span>",
        ],
    )
    number = _parse_float(direct)
    if number is not None:
        return number

    for pattern in (
        r'"newRating"\s*:\s*([0-9]{2,4})',
        r"\bnewRating\s*:\s*([0-9]{2,4})",
    ):
        match = re.search(pattern, text)
        if not match:
            continue
        raw = _parse_float(match.group(1))
        if raw is None:
            continue
        if raw > 100:
            return raw / 10
        return raw
    return None


def _extract_rating_count(text: str) -> str | None:
    for pattern in (
        r'"newRatingCount"\s*:\s*([0-9]+)',
        r"\bnewRatingCount\s*:\s*([0-9]+)",
        r'"ratingCount"\s*:\s*([0-9]+)',
        r"\bratingCount\s*:\s*([0-9]+)",
        r'rating_number_people[^>]*>\s*([0-9]+(?:\.[0-9]+)?)\s*</span>\s*人点评',
        r"book_rating_item_detail_count[\s\S]{0,300}?([0-9]+(?:\.[0-9]+)?(?:万|千)?)\s*人点评",
        r"([0-9]+(?:\.[0-9]+)?(?:万|千)?)\s*人点评",
    ):
        match = re.search(pattern, text)
        if match:
            normalized = _normalize_comment_count(match.group(1))
            if normalized:
                return normalized
    return None


def _normalize_comment_count(value: Any) -> str | None:
    text = _clean_text(str(value or ""))
    text = re.sub(r"\s+", "", text)
    if not text or not re.search(r"\d", text):
        return None
    match = re.search(r"([0-9]+(?:\.[0-9]+)?)(万|千)?(?:人)?(?:点评|评论|评价)?", text)
    if not match:
        return None
    count = f"{match.group(1)}{match.group(2) or ''}"
    return f"{count}人点评"


def _decode_json_string(value: str) -> str:
    try:
        return json.loads(f'"{value}"')
    except json.JSONDecodeError:
        return value


def _extract_stat(text: str, label: str) -> str | None:
    pattern = (
        rf"{re.escape(label)}[\s\S]{{0,260}}?"
        r'(?:content_number|subTitle_number)[^>]*>([0-9.]+)</span><span[^>]*>([^<]+)</span>'
    )
    match = re.search(pattern, text)
    if match:
        return f"{match.group(1)}{_clean_text(match.group(2))}"
    return None


def _format_percent(value: Any) -> str:
    number = _parse_float(value)
    return "" if number is None else f"{number:.1f}%"


def _status_label(status: str) -> str:
    return {
        "matched": "已匹配",
        "partial": "部分匹配",
        "not_found": "未找到",
        "failed": "抓取失败",
    }.get(status, status or "")


def _col_name(index: int) -> str:
    result = ""
    while index:
        index, rem = divmod(index - 1, 26)
        result = chr(65 + rem) + result
    return result


def safe_export_filename(prefix: str = "微信读书推荐值与项目匹配") -> str:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"{prefix}_{stamp}.xlsx"


def _content_types_xml() -> str:
    return """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
  <Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
  <Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>
  <Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>
  <Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>
  <Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>
</Types>"""


def _root_rels_xml() -> str:
    return """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
  <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>
  <Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties" Target="docProps/app.xml"/>
</Relationships>"""


def _workbook_xml() -> str:
    return """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <sheets><sheet name="推荐值与项目匹配" sheetId="1" r:id="rId1"/></sheets>
</workbook>"""


def _workbook_rels_xml() -> str:
    return """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
  <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
  <Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/>
</Relationships>"""


def _styles_xml() -> str:
    return """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <fonts count="1"><font><sz val="11"/><name val="Arial"/></font></fonts>
  <fills count="1"><fill><patternFill patternType="none"/></fill></fills>
  <borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>
  <cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>
  <cellXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/></cellXfs>
  <cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>
</styleSheet>"""


def _core_xml() -> str:
    now = datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ")
    return f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
  <dc:creator>AI ScriptBook</dc:creator>
  <cp:lastModifiedBy>AI ScriptBook</cp:lastModifiedBy>
  <dcterms:created xsi:type="dcterms:W3CDTF">{now}</dcterms:created>
  <dcterms:modified xsi:type="dcterms:W3CDTF">{now}</dcterms:modified>
</cp:coreProperties>"""


def _app_xml() -> str:
    return """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">
  <Application>AI ScriptBook</Application>
</Properties>"""
