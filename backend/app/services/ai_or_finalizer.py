from __future__ import annotations

import fcntl
import hashlib
import json
import logging
import re
import time
import zipfile
from contextvars import ContextVar
from datetime import datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path
from zoneinfo import ZoneInfo
from xml.etree import ElementTree as ET

import httpx
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from app.core.config import settings
from app.services.aigc_quality_gate import AigcManualReviewRequired
from app.services.deepseek import get_deepseek_client
from app.services.general_longform_runtime import atomic_json, atomic_text
from app.utils import paid_api_guard
from app.utils.distributed_limiter import distributed_semaphore

logger = logging.getLogger(__name__)
PROMPTS = Path(__file__).resolve().parents[1] / "prompts"
VERSION = "ai-or-final-v1"
_profile = ContextVar("ai_or_editor_profile", default="")
BASE = "https://api.ai-or.com/c/api/open/workflow"
WORD = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
PACKAGING_MARKERS = {"@最终书名", "@最终简介", "@目录", "@导言", "@正文"}
ET.register_namespace("w", WORD)


class FinalizerHold(AigcManualReviewRequired):
    pass


TRUNCATED = "双稿纠错输出不完整；已暂停，不自动重放付费调用"


class CorrectionTruncated(FinalizerHold):
    """The paid correction response arrived but stopped at the output limit."""


def enabled_for(created_at: datetime | None) -> bool:
    if not settings.AI_OR_FINALIZER_ENABLED:
        return False
    start = settings.AI_OR_FINALIZER_START_AT
    if not start:
        return True
    if created_at is None:
        return False
    boundary = datetime.fromisoformat(start.replace("Z", "+00:00"))
    return created_at.replace(tzinfo=created_at.tzinfo or timezone.utc) >= boundary


def digest(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def split_chunks(text: str, limit: int) -> list[str]:
    if not 1000 <= limit <= 50000:
        raise FinalizerHold("降AI分段上限必须为1000至50000字符")
    parts: list[str] = []
    buffer = ""
    metadata_section = False
    for line in text.splitlines(keepends=True):
        if _profile.get() == "editor-v0.5" and line.strip() in PACKAGING_MARKERS:
            metadata_section = line.strip() in {"@最终书名", "@最终简介", "@目录"}
        protected_metadata = _profile.get() == "editor-v0.5" and (metadata_section or line.strip() in PACKAGING_MARKERS)
        if line.lstrip().startswith(("#", "![")) or protected_metadata:
            if buffer:
                parts.append(buffer)
                buffer = ""
            parts.append(line)
            continue
        if len(line) > limit:
            raise FinalizerHold("存在超长单段，需人工分段；尚未调用降AI服务")
        if buffer and len(buffer) + len(line) > limit:
            parts.append(buffer)
            buffer = ""
        buffer += line
    if buffer:
        parts.append(buffer)
    return parts


def correction_pieces(text: str, limit: int) -> list[str]:
    """Split one rewritten chunk into smaller review windows; joining them restores the chunk."""
    if not 500 <= limit <= 50000:
        raise FinalizerHold("纠错子段上限必须为500至50000字符")
    parts: list[str] = []
    buffer = ""
    for line in text.splitlines(keepends=True):
        while len(line) > limit:
            if buffer:
                parts.append(buffer)
                buffer = ""
            boundaries = list(re.finditer(r"[。！？；.!?]", line[:limit]))
            boundary = boundaries[-1].end() if boundaries else limit
            parts.append(line[:boundary])
            line = line[boundary:]
        if buffer and len(buffer) + len(line) > limit:
            parts.append(buffer)
            buffer = ""
        buffer += line
    if buffer:
        parts.append(buffer)
    return parts


def structure(text: str) -> list[str]:
    return [line.strip() if line.strip().startswith(("#", "![")) else line.strip()[0]
            for line in text.splitlines() if line.strip().startswith(("#", "@", "！", "!["))]


def validate_rewrite(original: str, rewritten: str) -> None:
    if not rewritten.strip() or len(rewritten) < len(original) * 0.7:
        raise FinalizerHold("降AI结果为空或疑似截断，已保存结果，等待人工确认")
    if structure(original) != structure(rewritten):
        raise FinalizerHold("降AI改变了标题、图像或加工标记段落，已保存结果，等待人工确认")


def parse_issues(raw: str, text: str) -> tuple[str, list[dict]]:
    raw = raw.strip()
    if raw.startswith("```"):
        raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw)
    value = json.loads(raw)
    issues = value.get("issues")
    if not isinstance(issues, list):
        raise FinalizerHold("纠错模型未返回逐项纠错记录")
    spans = []
    for issue in issues:
        if not isinstance(issue, dict):
            raise FinalizerHold("纠错记录格式错误")
        before, after = issue.get("original"), issue.get("replacement")
        remark = issue.get("remark")
        if before == after:
            # 模型有时会把“无需改动”的句子也列进来；这是空操作，跳过即可，
            # 不能让一条 no-op 记录把整份交付停在 manual_review。
            continue
        if (not isinstance(before, str) or not before or not isinstance(after, str)
                or not isinstance(remark, str)
                or max(len(before), len(after), len(remark)) > 2000
                or text.count(before) != 1):
            raise FinalizerHold("纠错记录无法唯一定位或超出可读表格长度，已保存模型输出")
        # 模型偶尔漏写 severe 或写成字符串；缺省按“非严重”处理，避免一条格式瑕疵
        # 让整份交付停在 manual_review。
        severe = issue.get("severe")
        if isinstance(severe, str):
            severe = severe.strip().lower() in {"true", "1", "yes", "y", "是", "严重", "严重问题"}
        if not isinstance(severe, bool):
            severe = False
        issue = {"original": before, "replacement": after, "remark": remark, "severe": severe}
        position = text.index(before)
        spans.append((position, position + len(before), after, issue))
    spans.sort(key=lambda item: item[0])
    if any(left[1] > right[0] for left, right in zip(spans, spans[1:])):
        raise FinalizerHold("纠错记录重叠，不能安全自动应用")
    final = text
    for start, end, replacement, _issue in reversed(spans):
        final = final[:start] + replacement + final[end:]
    if structure(final) != structure(text):
        if _profile.get() != "editor-v0.5":
            raise FinalizerHold("纠错建议改变加工符号或标题，需要人工确认")
        original_headings = [line for line in text.splitlines() if line.lstrip().startswith("#")]
        final_headings = [line for line in final.splitlines() if line.lstrip().startswith("#")]
        original_images = [line for line in text.splitlines() if line.lstrip().startswith("![")]
        final_images = [line for line in final.splitlines() if line.lstrip().startswith("![")]
        if len(original_headings) != len(final_headings) or original_images != final_images:
            raise FinalizerHold("纠错不能增删章节或修改图像链接")
    return final, [item[3] for item in spans]


def _rewrite_body(text: str) -> dict:
    """官方降重接口 /workflow/text 的请求体。

    workModeId 与 workMode 二选一；实测两者都不传会返回 500「平台不能为空」，
    所以未配置时直接判为配置缺失，避免把付费调用浪费在必失败的请求上。
    """
    body: dict = {"accessKey": settings.AI_OR_ACCESS_KEY,
                  "accessSecret": settings.AI_OR_ACCESS_SECRET,
                  "inputText": text}
    if int(getattr(settings, "AI_OR_WORK_MODE_ID", 0) or 0):
        body["workModeId"] = int(settings.AI_OR_WORK_MODE_ID)
    elif str(getattr(settings, "AI_OR_WORK_MODE", "") or "").strip():
        body["workMode"] = str(settings.AI_OR_WORK_MODE).strip()
    else:
        raise FinalizerHold("未配置降AI工作模式（AI_OR_WORK_MODE_ID / AI_OR_WORK_MODE），原稿已保留")
    return body


def call_rewrite(text: str) -> dict:
    if not settings.AI_OR_ACCESS_KEY or not settings.AI_OR_ACCESS_SECRET:
        raise FinalizerHold("未配置降AI服务密钥，原稿已保留")
    from app.services.generation_spend import reserve

    if _profile.get() == "editor-v0.5":
        body = {"accessKey": settings.AI_OR_ACCESS_KEY, "accessSecret": settings.AI_OR_ACCESS_SECRET,
                "inputText": text, "repair_paragraphs": False}
        path = "/zhuque-direct"
    else:
        body = _rewrite_body(text)
        path = str(getattr(settings, "AI_OR_REWRITE_PATH", "/text") or "/text")
    if path not in {"/text", "/zhuque-direct"}:
        raise FinalizerHold("降AI接口路径不在已验证的端点范围内")
    reservation_index = reserve("ai-or-rewrite", max(0.01, len(text) * 0.2 / 1000 * 2) * 2, path)
    try:
        with httpx.Client(timeout=httpx.Timeout(settings.AI_OR_TIMEOUT_SECONDS, connect=30)) as client:
            response = client.post(BASE + path, headers={
                "X-Aigc-Platform": "api", "X-Aigc-Site": "www",
            }, json=body)
            response.raise_for_status()
            payload = response.json()
        result = payload.get("data", payload)
        if not isinstance(result, dict) or result.get("status") != 1:
            raise FinalizerHold("降AI接口未返回成功状态；请查看平台使用记录，禁止自动重试")
        if not isinstance(result.get("outputText"), str):
            raise FinalizerHold("降AI接口没有返回正文")
        saved = {key: result.get(key) for key in (
            "taskId", "status", "outputText", "consumeAmount", "consumeCharCount", "executeTime")}
        saved["budget_reservation_index"] = reservation_index
        saved["model"] = "zhuque-v1" if path == "/zhuque-direct" else str(body.get("workModeId") or body.get("workMode"))
        return saved
    except FinalizerHold:
        raise
    except Exception as exc:
        raise FinalizerHold("降AI请求结果不确定，已暂停避免重复扣费；请核对平台记录") from exc


def call_correction(original: str, rewritten: str, max_tokens: int | None = None) -> dict:
    rules_path = PROMPTS / ("editor_v05/correction.md" if _profile.get() == "editor-v0.5" else "ai_or_dual_manuscript_correction.md")
    rules = rules_path.read_text(encoding="utf-8")
    transport_path = PROMPTS / ("editor_v05/correction_transport.md" if _profile.get() == "editor-v0.5" else "ai_or_correction_transport.md")
    transport = transport_path.read_text(encoding="utf-8")
    prompt = rules + "\n\n" + transport + "\n\n" + json.dumps(
        {"原稿": original, "降AI后": rewritten}, ensure_ascii=False)
    model = settings.AI_OR_CORRECTION_MODEL or settings.DEEPSEEK_MODEL
    cap = int(max_tokens or settings.AI_OR_CORRECTION_MAX_TOKENS)
    from app.services.generation_spend import reserve_deepseek

    cap, reservation_index = reserve_deepseek(prompt, None, model, cap, with_index=True)
    with distributed_semaphore("deepseek-format", limit=settings.DEEPSEEK_FORMAT_MAX_CONCURRENT_CALLS):
        paid_api_guard.guard("deepseek")
        result = get_deepseek_client().with_options(max_retries=0).chat.completions.create(
            model=model, messages=[{"role": "user", "content": prompt}],
            temperature=0.1, response_format={"type": "json_object"},
            max_tokens=cap, extra_body={"thinking": {"type": "disabled"}},
        )
    choice = result.choices[0]
    if choice.finish_reason == "length":
        raise CorrectionTruncated(TRUNCATED)
    if choice.finish_reason != "stop":
        raise FinalizerHold(TRUNCATED)
    return {"text": choice.message.content or "", "model": model,
            "usage": result.usage.model_dump() if result.usage else None, "cost": None,
            "budget_reservation_index": reservation_index, "response_id": getattr(result, "id", None),
            "response_model": getattr(result, "model", None)}


def word_document(before: str, after: str | None = None) -> bytes:
    from io import BytesIO

    root = ET.Element(f"{{{WORD}}}document")
    body = ET.SubElement(root, f"{{{WORD}}}body")
    stamp = datetime.now(timezone.utc).isoformat()
    revision_id = 0

    def paragraph(text: str, revision: str | None = None) -> None:
        nonlocal revision_id
        node = ET.SubElement(body, f"{{{WORD}}}p")
        container = node
        if revision:
            revision_id += 1
            attrs = {f"{{{WORD}}}id": str(revision_id), f"{{{WORD}}}author": "降AI后纠错",
                     f"{{{WORD}}}date": stamp}
            properties = ET.SubElement(node, f"{{{WORD}}}pPr")
            mark = ET.SubElement(properties, f"{{{WORD}}}rPr")
            ET.SubElement(mark, f"{{{WORD}}}{revision}", attrs)
            container = ET.SubElement(node, f"{{{WORD}}}{revision}", attrs)
        run = ET.SubElement(container, f"{{{WORD}}}r")
        run_properties = ET.SubElement(run, f"{{{WORD}}}rPr")
        ET.SubElement(run_properties, f"{{{WORD}}}rFonts", {
            f"{{{WORD}}}ascii": "Arial Unicode MS", f"{{{WORD}}}hAnsi": "Arial Unicode MS",
            f"{{{WORD}}}eastAsia": "Arial Unicode MS"})
        child = ET.SubElement(run, f"{{{WORD}}}{'delText' if revision == 'del' else 't'}")
        child.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        child.text = text

    previous = before.split("\n")
    if after is None:
        for line in previous:
            paragraph(line)
    else:
        current = after.split("\n")
        for action, left_start, left_end, right_start, right_end in SequenceMatcher(
                None, previous, current, autojunk=False).get_opcodes():
            if action == "equal":
                for line in previous[left_start:left_end]:
                    paragraph(line)
            else:
                for line in previous[left_start:left_end]:
                    paragraph(line, "del")
                for line in current[right_start:right_end]:
                    paragraph(line, "ins")
    settings_xml = f'<w:settings xmlns:w="{WORD}"><w:trackRevisions/></w:settings>'
    content_types = '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/><Override PartName="/word/settings.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.settings+xml"/></Types>'
    root_rels = '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>'
    document_rels = '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/settings" Target="settings.xml"/></Relationships>'
    stream = BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in {"[Content_Types].xml": content_types, "_rels/.rels": root_rels,
                              "word/document.xml": ET.tostring(root, encoding="utf-8", xml_declaration=True),
                              "word/settings.xml": settings_xml, "word/_rels/document.xml.rels": document_rels}.items():
            archive.writestr(name, content)
    return stream.getvalue()


def write_delivery(directory: Path, original: str, rewritten: str, final: str,
                   issues: list[dict], title: str, stamp: str) -> dict:
    replacements = str.maketrans('/\\:*?"<>|', '／＼：＊？＂＜＞｜')
    safe_title = re.sub(r"\s+", "_", title.strip()).translate(replacements) or "书稿"
    stem = f"【降AI后纠错】_{safe_title}_{stamp}"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "双稿纠错"
    sheet.append(["原文是什么", "建议改成什么", "备注"])
    for issue in issues:
        remark = issue["remark"]
        if issue["severe"] and "严重问题" not in remark:
            remark = "严重问题；" + remark
        sheet.append([issue["original"], issue["replacement"], remark])
        if issue["severe"]:
            sheet.cell(sheet.max_row, 1).font = Font(name="微软雅黑", size=10, color="FF0000")
    if not issues:
        sheet.append(["未发现需要修改的明显错误或局部表达问题，文稿可以进入电子书制作。", "", ""])
    border = Border(left=Side(style="thin"), right=Side(style="thin"),
                    top=Side(style="thin"), bottom=Side(style="thin"))
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            color = "FF0000" if cell.font.color and cell.font.color.type == "rgb" and cell.font.color.rgb == "00FF0000" else None
            cell.font = Font(name="微软雅黑", size=10, bold=cell.row == 1, color=color)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = border
        sheet.row_dimensions[row[0].row].height = min(409, max(30, max(
            len(str(cell.value or "")) / (28 if cell.column < 3 else 16) for cell in row) * 15 + 15))
    for column, width in [("A", 60), ("B", 60), ("C", 36)]:
        sheet.column_dimensions[column].width = width
    workbook.save(directory / (stem + ".xlsx"))
    (directory / (stem + ".docx")).write_bytes(word_document(rewritten, final))
    (directory / "原始文稿.docx").write_bytes(word_document(original))
    atomic_text(directory / "原始文稿.md", original)
    atomic_text(directory / "降AI后原稿.md", rewritten)
    atomic_text(directory / "final.md", final)
    filenames = [stem + ".xlsx", stem + ".docx", "原始文稿.docx", "原始文稿.md", "降AI后原稿.md", "final.md"]
    with zipfile.ZipFile(directory / "delivery.zip", "w", zipfile.ZIP_DEFLATED) as archive:
        for name in filenames:
            archive.write(directory / name, name)
    return {"docx": stem + ".docx", "xlsx": stem + ".xlsx", "zip": "delivery.zip"}


def finalize_manuscript(text: str, output_dir: Path, *, title: str = "书稿", progress=None, profile: str = "") -> str:
    token = _profile.set(profile)
    try:
        return _finalize_manuscript(text, output_dir, title=title, progress=progress)
    finally:
        _profile.reset(token)


def _finalize_manuscript(text: str, output_dir: Path, *, title: str = "书稿", progress=None) -> str:
    directory = output_dir / "ai_or_finalizer"
    directory.mkdir(parents=True, exist_ok=True)
    with (directory / "workflow.lock").open("a") as lock:
        try:
            fcntl.flock(lock, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError as exc:
            raise FinalizerHold("该书降AI交付步骤正在运行，禁止并行重复扣费") from exc
        manifest_path = directory / "manifest.json"
        state = json.loads(manifest_path.read_text()) if manifest_path.exists() else {
            "version": _profile.get() or VERSION, "source_hash": digest(text), "status": "running",
            "stamp": datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y年%-m月%-d日_%H时%M分%S秒"),
            "chunks": [], "detection": "not_requested"}
        if state["source_hash"] != digest(text) or state.get("version") != (_profile.get() or VERSION):
            raise FinalizerHold("原稿或处理版本已变更，已有付费断点不能覆盖；请人工确认")
        if _profile.get() == "editor-v0.5":
            prompt_hash = digest((PROMPTS / "editor_v05/correction.md").read_text() +
                                 (PROMPTS / "editor_v05/correction_transport.md").read_text())
            if state.setdefault("correction_prompt_hash", prompt_hash) != prompt_hash:
                raise FinalizerHold("纠错提示词已变更，不能覆盖已有付费断点")
        if state["status"] == "completed":
            final = (directory / "final.md").read_text(encoding="utf-8")
            if digest(final) != state["final_hash"]:
                raise FinalizerHold("最终交付文件校验失败")
            return final
        atomic_text(directory / "source.md", text)
        atomic_json(manifest_path, state)
        chunks = split_chunks(text, settings.AI_OR_CHUNK_CHARS)
        rewrites, finals, issues = [], [], []
        metadata_section = False
        try:
            for index, chunk in enumerate(chunks):
                if _profile.get() == "editor-v0.5" and chunk.strip() in PACKAGING_MARKERS:
                    metadata_section = chunk.strip() in {"@最终书名", "@最终简介", "@目录"}
                protected_metadata = _profile.get() == "editor-v0.5" and (metadata_section or chunk.strip() in PACKAGING_MARKERS)
                if not chunk.strip() or chunk.lstrip().startswith(("#", "![")) or protected_metadata:
                    rewrites.append(chunk)
                    finals.append(chunk)
                    if len(state["chunks"]) <= index:
                        state["chunks"].append({"hash": digest(chunk), "actions": {}, "protected": True})
                    continue
                if progress:
                    progress(f"降AI及双稿纠错 {index + 1}/{len(chunks)}")
                if len(state["chunks"]) <= index:
                    state["chunks"].append({"hash": digest(chunk), "actions": {}})
                entry = state["chunks"][index]
                if entry["hash"] != digest(chunk):
                    raise FinalizerHold("分段断点与原稿不一致")
                def paid(key: str, artifact: Path, call, cap: int | None = None) -> dict:
                    if artifact.exists():
                        result = json.loads(artifact.read_text())
                        entry["actions"][key] = "saved"
                        atomic_json(manifest_path, state)
                        return result
                    status = entry["actions"].get(key)
                    if status == "inflight":
                        raise FinalizerHold("上次付费请求结果不确定；请核对平台记录后人工恢复")
                    if isinstance(status, str) and status.startswith("truncated:") and (
                            cap is None or cap <= int(status.split(":", 1)[1])):
                        raise FinalizerHold("纠错子段输出仍超过上限；请调大 AI_OR_CORRECTION_MAX_TOKENS "
                                            "或调小 AI_OR_CORRECTION_CHUNK_CHARS 后人工恢复")
                    entry["actions"][key] = "inflight"
                    atomic_json(manifest_path, state)
                    started = time.monotonic()
                    try:
                        result = call()
                    except CorrectionTruncated:
                        entry["actions"][key] = f"truncated:{cap}"
                        atomic_json(manifest_path, state)
                        raise
                    except AigcManualReviewRequired as exc:
                        from app.services.generation_spend import GenerationBudgetExceeded

                        if isinstance(exc, GenerationBudgetExceeded):
                            entry["actions"].pop(key, None)
                            atomic_json(manifest_path, state)
                        raise
                    atomic_json(artifact, result)
                    logger.info("finalizer action=%s model=%s cost=%s status=completed elapsed=%.2f",
                                key, result.get("model", "zhuque-v1-paragraph-enhanced"),
                                result.get("consumeAmount", result.get("cost")), time.monotonic() - started)
                    entry["actions"][key] = "saved"
                    atomic_json(manifest_path, state)
                    return result

                result = paid("rewrite", directory / f"{index:05d}_rewrite.json", lambda: call_rewrite(chunk))
                from app.services.generation_spend import settle_rewrite

                settle_rewrite(result.get("budget_reservation_index"), directory / f"{index:05d}_rewrite.json")
                ending = chunk[len(chunk.rstrip("\n")):]
                rewritten = result["outputText"].rstrip("\n") + ending
                validate_rewrite(chunk, rewritten)

                legacy = directory / f"{index:05d}_correction.json"
                if legacy.exists():
                    # Whole-chunk correction saved by the earlier version: reuse it, never pay again.
                    entry["actions"]["correction"] = "saved"
                    corrected, chunk_issues = parse_issues(json.loads(legacy.read_text())["text"], rewritten)
                else:
                    if entry["actions"].get("correction") == "inflight":
                        # The earlier whole-chunk call only leaves this state with the truncation error
                        # when its response actually arrived and stopped at the output limit.
                        if state.get("error") != TRUNCATED:
                            raise FinalizerHold("上次付费请求结果不确定；请核对平台记录后人工恢复")
                        entry["actions"]["correction"] = "truncated:4096"
                        atomic_json(manifest_path, state)
                    limit = settings.AI_OR_CORRECTION_CHUNK_CHARS
                    cap = settings.AI_OR_CORRECTION_MAX_TOKENS
                    pieces = correction_pieces(rewritten, limit)
                    plan = {"limit": limit, "hashes": [digest(piece) for piece in pieces]}
                    if entry.setdefault("correction_split", plan) != plan:
                        raise FinalizerHold("纠错子段断点与当前分段设置不一致；请恢复原分段设置后再续跑")
                    atomic_json(manifest_path, state)
                    corrected_parts, chunk_issues = [], []
                    for part, piece in enumerate(pieces):
                        if not piece.strip():
                            corrected_parts.append(piece)
                            continue
                        result = paid(f"correction_{part}", directory / f"{index:05d}_correction_{part:02d}.json",
                                      lambda piece=piece: call_correction(chunk, piece, cap), cap)
                        from app.services.generation_spend import settle_deepseek_usage

                        settle_deepseek_usage(result.get("budget_reservation_index"), directory / f"{index:05d}_correction_{part:02d}.json")
                        piece_final, piece_issues = parse_issues(result["text"], piece)
                        corrected_parts.append(piece_final)
                        chunk_issues.extend(piece_issues)
                    corrected = "".join(corrected_parts)
                rewrites.append(rewritten)
                finals.append(corrected)
                issues.extend(chunk_issues)
            rewritten = "".join(rewrites)
            final = "".join(finals)
            if any(issue["severe"] for issue in issues):
                atomic_json(directory / "issues.json", issues)
            state["stamp"] = datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y年%-m月%-d日_%H时%M分%S秒")
            artifacts = write_delivery(directory, text, rewritten, final, issues, title, state["stamp"])
            severe_count = sum(1 for issue in issues if issue["severe"])
            if severe_count and bool(getattr(settings, "AI_OR_SEVERE_BLOCKS_DELIVERY", False)):
                raise FinalizerHold("双稿纠错发现严重问题，已生成标红Excel和修订Word，需人工确认后交付")
            state.update(status="completed", final_hash=digest(final), artifacts=artifacts,
                         issue_count=len(issues), severe_issue_count=severe_count,
                         correction_model=settings.AI_OR_CORRECTION_MODEL or settings.DEEPSEEK_MODEL)
            atomic_json(manifest_path, state)
            return final
        except Exception as exc:
            state.update(status="manual_review", error=str(exc) if isinstance(exc, AigcManualReviewRequired) else type(exc).__name__)
            atomic_json(manifest_path, state)
            logger.error("finalizer status=manual_review error_type=%s", type(exc).__name__)
            if isinstance(exc, AigcManualReviewRequired):
                raise
            raise FinalizerHold("降AI或双稿纠错未完成，原稿与断点已保存，需人工核对") from exc


def finalize_task_manuscript(task, text: str, output_dir: Path, *, progress=None) -> str:
    if not enabled_for(task.created_at):
        return text
    from app.services.editor_v05 import applies, finalize_editor_task

    if applies(task):
        return finalize_editor_task(task, text, output_dir, progress=progress)
    cfg = dict(task.config or {})
    heading = next((line.lstrip("#").strip() for line in text.splitlines() if line.startswith("#")), "书稿")
    final = finalize_manuscript(text, output_dir, title=str(cfg.get("title") or heading), progress=progress)
    manifest_path = output_dir / "ai_or_finalizer" / "manifest.json"
    severe_count = 0
    if manifest_path.exists():
        try:
            severe_count = int(json.loads(manifest_path.read_text()) .get("severe_issue_count") or 0)
        except (OSError, ValueError):
            severe_count = 0
    cfg.update(ai_or_finalizer_status="completed", ai_or_detection_status="not_requested",
               ai_or_severe_issue_count=severe_count,
               ai_or_correction_model=settings.AI_OR_CORRECTION_MODEL or settings.DEEPSEEK_MODEL)
    task.config = cfg
    return final
