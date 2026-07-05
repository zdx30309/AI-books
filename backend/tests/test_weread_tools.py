from __future__ import annotations

import asyncio
import base64
import json
import zipfile
from io import BytesIO

from openpyxl import Workbook

from app.api.routers.weread import WeReadBookRow, WeReadExportRequest
from app.services import weread_tools
from app.services.weread_tools import (
    CatalogBook,
    WeReadBookResult,
    build_weread_xlsx,
    catalog_books_to_text,
    parse_catalog,
    parse_catalog_workbook,
)


def _xlsx_shared_text(content: bytes) -> str:
    with zipfile.ZipFile(BytesIO(content)) as zf:
        return zf.read("xl/sharedStrings.xml").decode("utf-8")


def test_parse_catalog_keeps_weread_link() -> None:
    rows = parse_catalog("https://weread.qq.com/web/reader/aea32010813ab9e28g016851")

    assert len(rows) == 1
    assert rows[0].source_url == "https://weread.qq.com/web/reader/aea32010813ab9e28g016851"


def test_parse_catalog_uses_book_title_from_copied_table_rows(monkeypatch) -> None:
    monkeypatch.setattr(
        weread_tools,
        "_PROJECT_INDEX_CACHE",
        [{"editor": "单华伟", "_title_key": "", "_url_key": ""}],
    )

    rows = parse_catalog("编辑名\t书名\n单华伟\t肉蒲团（白话版）")

    assert [row.title for row in rows] == ["肉蒲团（白话版）"]


def test_parse_catalog_skips_editor_name_only_lines(monkeypatch) -> None:
    monkeypatch.setattr(
        weread_tools,
        "_PROJECT_INDEX_CACHE",
        [{"editor": "单华伟", "_title_key": "", "_url_key": ""}],
    )

    rows = parse_catalog("单华伟\n肉蒲团（白话版）")

    assert [row.title for row in rows] == ["肉蒲团（白话版）"]


def test_parse_catalog_workbook_supports_project_progress_excel(monkeypatch) -> None:
    monkeypatch.setattr(
        weread_tools,
        "_PROJECT_INDEX_CACHE",
        [{"editor": "单华伟", "_title_key": "", "_url_key": ""}],
    )
    workbook = Workbook()
    ws = workbook.active
    ws.title = "已完成书目"
    ws.append(["编辑名", "书名", "原著作者", "微信读书链接", "推荐值", "阅读"])
    ws.append(["单华伟", "论老年 论友谊", "西塞罗", "https://weread.qq.com/web/reader/book001", 0.8, 623])
    ws.append([None, "论责任", "西塞罗", "https://weread.qq.com/web/reader/book002", 0.9, 176])
    ws.append([None, "单华伟", None, None, None, None])
    month = workbook.create_sheet("202606")
    month.append(["索引字段（导出时备份用，勿动）", "编辑", "书名", "简介", "父记录"])
    month.append(["单华伟", "单华伟", "职业网球到底多烧钱？", "简介", None])
    month.append(["单华伟", "单华伟", "论责任", "重复", None])

    output = BytesIO()
    workbook.save(output)

    rows = parse_catalog_workbook(output.getvalue())

    assert [(row.title, row.author, row.source_url) for row in rows] == [
        ("论老年 论友谊", "西塞罗", "https://weread.qq.com/web/reader/book001"),
        ("论责任", "西塞罗", "https://weread.qq.com/web/reader/book002"),
        ("职业网球到底多烧钱？", None, None),
    ]


def test_parse_catalog_workbook_supports_weread_export_headers() -> None:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "推荐值与项目匹配"
    ws.append(
        [
            "原始书名",
            "原始链接",
            "匹配书名",
            "作者",
            "译者",
            "版权方（公司名字）",
            "推荐值",
            "阅读人数",
            "评论数量",
        ]
    )
    ws.append(["多余的话", None, "多余的话", "瞿秋白", None, "明天远航", None, None, None])
    ws.append(
        [
            "",
            "https://weread.qq.com/web/reader/book002",
            "段永平投资问答录：大道至简的财富逻辑",
            "马小奔",
            None,
            "明天远航",
            "79.2%",
            "688人",
            "44人点评",
        ]
    )

    output = BytesIO()
    workbook.save(output)

    rows = parse_catalog_workbook(output.getvalue())

    assert [(row.title, row.author, row.source_url) for row in rows] == [
        ("多余的话", "瞿秋白", None),
        (
            "段永平投资问答录：大道至简的财富逻辑",
            "马小奔",
            "https://weread.qq.com/web/reader/book002",
        ),
    ]


def test_catalog_books_to_text_keeps_excel_fields_for_chunked_fetch() -> None:
    text = catalog_books_to_text(
        [
            CatalogBook("论责任", "西塞罗", "https://weread.qq.com/web/reader/book002"),
            CatalogBook("职业网球到底多烧钱？", None, None),
        ]
    )

    assert text.splitlines() == [
        "书名\t原著作者\t微信读书链接",
        "论责任\t西塞罗\thttps://weread.qq.com/web/reader/book002",
        "职业网球到底多烧钱？\t\t",
    ]


def test_project_index_match_adds_company_and_existing_excel_fields(monkeypatch) -> None:
    monkeypatch.setattr(
        weread_tools,
        "_PROJECT_INDEX_CACHE",
        [
            {
                "title": "黄金投资指南",
                "_title_key": weread_tools.normalize_title("黄金投资指南"),
                "project_weread_url": "https://weread.qq.com/web/reader/demo123",
                "_url_key": "reader:demo123",
                "editor": "单华伟",
                "company_name": "汇安",
                "project_recommendation": "0.805",
                "project_comment_count": "44",
                "project_reading_count": "623",
                "source_sheet": "已完成书目",
            }
        ],
    )
    result = WeReadBookResult(
        query_title="https://weread.qq.com/web/reader/demo123?from=share",
        query_author=None,
        source_url="https://weread.qq.com/web/reader/demo123?from=share",
        status="failed",
        error="network failed",
    )

    weread_tools._attach_project_match(result)

    assert result.company_name == "汇安"
    assert result.project_editor == "单华伟"
    assert result.project_match_status == "链接匹配"
    assert result.recommendation == 80.5
    assert result.reading_count == "623"
    assert result.comment_count == "44人点评"
    assert result.status == "matched"
    assert result.error is None


def test_extract_copyright_holder_from_weread_payload() -> None:
    body = '{"publisher":"广西师范大学出版社","copyrightOwner":"","copyrightInfo":{"id":1,"name":"简读互动"},"name":"肉蒲团"}'

    assert weread_tools._extract_copyright_holder(body) == "简读互动"


def test_extract_copyright_holder_from_book_detail_payload() -> None:
    body = 'bookInfo:{publisher:b,copyrightInfo:{id:u,name:"简读互动",userVid:907735518,role:l},title:n}'

    assert weread_tools._extract_copyright_holder(body) == "简读互动"


def test_extract_copyright_id_from_book_detail_payload() -> None:
    body = '{"bookInfo":{"cpid":39149647,"copyrightInfo":{"id":39149647,"name":"佳润互众"}}}'

    assert weread_tools._extract_copyright_id(body) == "39149647"


def test_parse_company_lookup_input_keeps_seed_link() -> None:
    lookup = weread_tools._parse_company_lookup_input(
        "佳润互众\nhttps://weread.qq.com/web/reader/b3a32f00813aba926g018a42"
    )

    assert lookup.company_name == "佳润互众"
    assert lookup.seed_items == [
        CatalogBook(
            title="https://weread.qq.com/web/reader/b3a32f00813aba926g018a42",
            source_url="https://weread.qq.com/web/reader/b3a32f00813aba926g018a42",
        )
    ]


def test_parse_company_lookup_input_extracts_ids() -> None:
    lookup = weread_tools._parse_company_lookup_input(
        "佳润互众\n"
        "https://weread.qq.com/web/search?cpId=39149647&subjectId=sub_123\n"
        "版权方ID：8848"
    )

    assert lookup.company_name == "佳润互众"
    assert lookup.copyright_ids == {"39149647", "8848"}
    assert lookup.subject_ids == {"sub_123"}


def test_company_index_upserts_and_returns_cached_cards(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(weread_tools, "WEREAD_COMPANY_INDEX_PATH", tmp_path / "weread_company_index.json")
    monkeypatch.setattr(weread_tools, "_COMPANY_INDEX_CACHE", None)

    weread_tools._upsert_company_index_from_result(
        WeReadBookResult(
            query_title="天涯神贴智慧大全",
            query_author="玄谷",
            matched_title="天涯神贴智慧大全",
            matched_author="玄谷",
            recommendation=66.4,
            rating_count="223人点评",
            reading_count="4.1万人",
            company_name="佳润互众",
            copyright_id="39149647",
            source_url="https://weread.qq.com/web/reader/b3a32f00813aba926g018a42",
            status="matched",
        )
    )

    cards = weread_tools._company_candidate_cards_from_cache("佳润互众", target_ids=set(), max_books=10)

    assert cards[0]["title"] == "天涯神贴智慧大全"
    assert cards[0]["company_name"] == "佳润互众"
    assert cards[0]["copyright_id"] == "39149647"
    assert cards[0]["reading_count"] == "4.1万人"


def test_parse_api_book_cards_keeps_company_metrics_and_url() -> None:
    cards = weread_tools._parse_api_book_cards(
        {
            "books": [
                {
                    "bookInfo": {
                        "bookId": "3300170628",
                        "encodeId": "b3a32f00813aba926g018a42",
                        "title": "天涯神贴智慧大全",
                        "author": "玄谷",
                        "newRating": 662,
                        "newRatingCount": 222,
                        "copyrightInfo": {"id": 39149647, "name": "佳润互众"},
                    },
                    "readingCount": 674,
                    "searchIdx": 2,
                }
            ]
        }
    )

    assert cards == [
        {
            "title": "天涯神贴智慧大全",
            "author": "玄谷",
            "source_url": "https://weread.qq.com/web/reader/b3a32f00813aba926g018a42",
            "company_name": "佳润互众",
            "copyright_id": "39149647",
            "recommendation": 66.2,
            "rating_count": "222人点评",
            "translator": None,
            "word_count": None,
            "today_reading": "674人",
            "reading_count": "674人",
            "search_idx": 2,
        }
    ]


def test_parse_weread_capture_cards_from_har_content() -> None:
    payload = {
        "books": [
            {
                "bookInfo": {
                    "encodeId": "book-a",
                    "title": "天涯神帖智慧大全",
                    "author": "玄谷",
                    "newRating": 664,
                    "newRatingCount": 223,
                    "copyrightInfo": {"id": 39149647, "name": "佳润互众"},
                },
                "readingCount": 100,
                "searchIdx": 5,
            }
        ],
        "hasMore": False,
    }
    har = {
        "log": {
            "entries": [
                {"response": {"content": {"text": json.dumps(payload, ensure_ascii=False)}}}
            ]
        }
    }

    cards = weread_tools._parse_weread_capture_book_cards(
        json.dumps(har, ensure_ascii=False).encode("utf-8"),
        filename="capture.har",
    )

    assert len(cards) == 1
    assert cards[0]["title"] == "天涯神帖智慧大全"
    assert cards[0]["company_name"] == "佳润互众"
    assert cards[0]["copyright_id"] == "39149647"


def test_parse_weread_capture_cards_decodes_base64_har_content() -> None:
    payload = {"totalResult": {"results": [{"bookInfo": {"encodeId": "book-a", "title": "孟姜女", "author": "张恨水"}}]}}
    encoded = base64.b64encode(json.dumps(payload, ensure_ascii=False).encode("utf-8")).decode("ascii")
    har = {"log": {"entries": [{"response": {"content": {"text": encoded, "encoding": "base64"}}}]}}

    cards = weread_tools._parse_weread_capture_book_cards(json.dumps(har).encode("utf-8"))

    assert [card["title"] for card in cards] == ["孟姜女"]


def test_parse_weread_capture_cards_from_store_search_har_pages() -> None:
    payload = {
        "sid": "search-session",
        "results": [
            {
                "title": "ebooks",
                "scopeCount": 2,
                "books": [
                    {
                        "bookInfo": {
                            "bookId": "book-1",
                            "title": "Book One",
                            "author": "Author A",
                            "newRating": 885,
                            "newRatingCount": 11,
                        },
                        "readingCount": 123,
                        "searchIdx": 1,
                    },
                    {
                        "bookInfo": {
                            "bookId": "book-2",
                            "title": "Book Two",
                            "author": "Author B",
                            "newRating": 800,
                            "newRatingCount": 22,
                        },
                        "readingCount": 456,
                        "searchIdx": 2,
                    },
                ],
            }
        ],
    }
    har = {
        "log": {
            "entries": [
                {"request": {"url": "https://weread.qq.com/"}, "response": {"content": {"text": ""}}},
                {
                    "request": {"url": "https://weread.qq.com/api/store/search?keyword=x&scope=17"},
                    "response": {"content": {"mimeType": "application/json", "text": json.dumps(payload)}},
                },
            ]
        }
    }

    cards = weread_tools._parse_weread_capture_book_cards(json.dumps(har).encode("utf-8"))

    assert [card["title"] for card in cards] == ["Book One", "Book Two"]
    assert cards[0]["reading_count"] == "123人"
    assert cards[1]["recommendation"] == 80.0


def test_parse_fanqie_capture_cards_from_har_content() -> None:
    payload = {
        "data": {
            "book_list": [
                {
                    "book_id": "123456",
                    "book_name": "Fanqie Book",
                    "author_name": "Fanqie Author",
                    "score": "9.6",
                    "read_count": "166.9万人读过",
                    "comment_count": 88,
                    "word_count": 120000,
                }
            ]
        }
    }
    har = {
        "log": {
            "entries": [
                {
                    "request": {"url": "https://fanqienovel.com/api/book/list"},
                    "response": {"content": {"mimeType": "application/json", "text": json.dumps(payload)}},
                }
            ]
        }
    }

    cards = weread_tools._parse_fanqie_capture_book_cards(json.dumps(har).encode("utf-8"))

    assert len(cards) == 1
    assert cards[0]["title"] == "Fanqie Book"
    assert cards[0]["author"] == "Fanqie Author"
    assert cards[0]["recommendation"] == 96.0
    assert cards[0]["reading_count"] == "166.9万人读过"
    assert cards[0]["rating_count"] == "88人点评"
    assert cards[0]["source_url"] == "https://fanqienovel.com/page/123456"
    assert cards[0]["platform"] == "番茄小说"


def test_analyze_fanqie_capture_report_returns_platform() -> None:
    payload = {
        "books": [
            {
                "bookId": "abc001",
                "bookName": "Second Book",
                "authorName": "Writer B",
                "rankScore": 88,
                "readCount": 2100,
            }
        ]
    }
    report = asyncio.run(
        weread_tools.analyze_fanqie_company_capture_report(
            "番茄样例",
            json.dumps(payload).encode("utf-8"),
            filename="fanqie.json",
            limit=10,
        )
    )

    assert report.rows
    assert report.rows[0].platform == "番茄小说"
    assert report.rows[0].matched_title == "Second Book"
    assert report.rows[0].source_url == "https://fanqienovel.com/page/abc001"


def test_fetch_app_publisher_book_cards_uses_store_search_pagination(monkeypatch) -> None:
    monkeypatch.setenv("WEREAD_COOKIE", "wr-test-cookie")
    calls = []

    class FakeResponse:
        status_code = 200

        def __init__(self, payload):
            self._payload = payload

        def raise_for_status(self):
            return None

        def json(self):
            return self._payload

    class FakeClient:
        async def get(self, url, *, params=None, headers=None):
            calls.append((url, dict(params or {}), dict(headers or {})))
            if (params or {}).get("maxIdx") == 0:
                return FakeResponse(
                    {
                        "books": [
                            {
                                "bookInfo": {
                                    "encodeId": "book-a",
                                    "title": "\u5929\u6daf\u795e\u8d34\u667a\u6167\u5927\u5168",
                                    "author": "\u7384\u8c37",
                                    "newRating": 664,
                                    "newRatingCount": 223,
                                },
                                "readingCount": 100,
                                "searchIdx": 5,
                            }
                        ],
                        "hasMore": True,
                    }
                )
            return FakeResponse(
                {
                    "books": [
                        {
                            "bookInfo": {
                                "encodeId": "book-b",
                                "title": "\u5b5f\u59dc\u5973\uff08\u7ecf\u5178\u7248\uff09",
                                "author": "\u5f20\u6068\u6c34",
                                "newRating": 950,
                                "newRatingCount": 34,
                            },
                            "readingCount": 200,
                            "searchIdx": 9,
                        }
                    ],
                    "hasMore": False,
                }
            )

    cards = asyncio.run(
        weread_tools._fetch_app_publisher_book_cards(
            FakeClient(),
            "\u4f73\u6da6\u4e92\u4f17",
            max_books=20,
            from_book_id="seed-book-id",
            copyright_ids={"39149647"},
        )
    )

    assert len(cards) == 2
    assert [call[1]["maxIdx"] for call in calls] == [0, 5]
    assert calls[0][1]["type"] == 8
    assert calls[0][1]["keyword"] == "\u4f73\u6da6\u4e92\u4f17"
    assert calls[0][1]["fromBookId"] == "seed-book-id"
    assert calls[0][1]["cpid"] == "39149647"
    assert calls[0][1]["recordType"] == 2
    assert calls[0][2]["Cookie"] == "wr-test-cookie"
    assert cards[0]["source_kind"] == "app_publisher"
    assert cards[0]["company_name"] == "\u4f73\u6da6\u4e92\u4f17"


def test_fetch_uncertified_opus_book_cards_uses_total_result_pagination(monkeypatch) -> None:
    monkeypatch.setenv("WEREAD_COOKIE", "wr-test-cookie")
    calls = []

    class FakeResponse:
        status_code = 200

        def __init__(self, payload):
            self._payload = payload

        def raise_for_status(self):
            return None

        def json(self):
            return self._payload

    class FakeClient:
        async def get(self, url, *, params=None, headers=None):
            calls.append((url, dict(params or {}), dict(headers or {})))
            if (params or {}).get("maxIdx") == 0:
                return FakeResponse(
                    {
                        "totalResult": {
                            "results": [
                                {
                                    "bookInfo": {
                                        "encodeId": "book-a",
                                        "title": "天涯神帖智慧大全",
                                        "author": "玄谷",
                                        "newRating": 664,
                                        "newRatingCount": 223,
                                    },
                                    "readingCount": 100,
                                    "searchIdx": 5,
                                }
                            ],
                            "hasMore": True,
                        }
                    }
                )
            return FakeResponse(
                {
                    "totalResult": {
                        "results": [
                            {
                                "bookInfo": {
                                    "encodeId": "book-b",
                                    "title": "孟姜女（经典版）",
                                    "author": "张恨水",
                                    "newRating": 950,
                                    "newRatingCount": 34,
                                },
                                "readingCount": 200,
                                "searchIdx": 9,
                            }
                        ],
                        "hasMore": False,
                    }
                }
            )

    cards = asyncio.run(
        weread_tools._fetch_uncertified_opus_book_cards(
            FakeClient(),
            "佳润互众",
            max_books=20,
            from_book_id="seed-book-id",
        )
    )

    assert len(cards) == 2
    assert [call[1]["maxIdx"] for call in calls] == [0, 5]
    assert calls[0][0].endswith("/shelf/uncertifiedopus")
    assert calls[0][1]["bookId"] == "seed-book-id"
    assert cards[0]["source_kind"] == "app_uncertified_opus"


def test_fetch_subject_book_cards_uses_subject_info_pagination(monkeypatch) -> None:
    monkeypatch.setenv("WEREAD_COOKIE", "wr-test-cookie")
    calls = []

    class FakeResponse:
        status_code = 200

        def __init__(self, payload):
            self._payload = payload

        def raise_for_status(self):
            return None

        def json(self):
            return self._payload

    class FakeClient:
        async def post(self, url, *, json=None, headers=None):
            calls.append((url, dict(json or {}), dict(headers or {})))
            if (json or {}).get("maxIdx") == 0:
                return FakeResponse(
                    {
                        "totalResult": {
                            "results": [
                                {
                                    "bookInfo": {
                                        "encodeId": "book-a",
                                        "title": "天涯神帖智慧大全",
                                        "author": "玄谷",
                                        "newRating": 664,
                                        "newRatingCount": 223,
                                    },
                                    "readingCount": 100,
                                    "searchIdx": 5,
                                }
                            ],
                            "hasMore": True,
                        }
                    }
                )
            return FakeResponse({"totalResult": {"results": [], "hasMore": False}})

    cards = asyncio.run(weread_tools._fetch_subject_book_cards(FakeClient(), "sub_123", max_books=20))

    assert len(cards) == 1
    assert calls[0][0].endswith("/special/subject/info")
    assert calls[0][1]["subjectId"] == "sub_123"
    assert calls[0][1]["maxIdx"] == 0
    assert cards[0]["source_kind"] == "app_subject"


def test_fetch_gateway_publisher_book_cards_uses_agent_gateway_pagination(monkeypatch) -> None:
    monkeypatch.setenv("WEREAD_API_KEY", "wrk-test-key")
    calls = []

    class FakeResponse:
        status_code = 200

        def __init__(self, payload):
            self._payload = payload

        def raise_for_status(self):
            return None

        def json(self):
            return self._payload

    class FakeClient:
        async def post(self, url, *, json=None, headers=None):
            calls.append((url, dict(json or {}), dict(headers or {})))
            if (json or {}).get("maxIdx") == 0:
                return FakeResponse(
                    {
                        "errcode": 0,
                        "data": {
                            "books": [
                                {
                                    "bookInfo": {
                                        "encodeId": "book-a",
                                        "title": "天涯神贴智慧大全",
                                        "author": "玄谷",
                                        "newRating": 664,
                                        "newRatingCount": 223,
                                    },
                                    "readingCount": 100,
                                    "searchIdx": 5,
                                }
                            ],
                            "hasMore": True,
                        },
                    }
                )
            return FakeResponse(
                {
                    "errcode": 0,
                    "data": {
                        "books": [
                            {
                                "bookInfo": {
                                    "encodeId": "book-b",
                                    "title": "孟姜女（经典版）",
                                    "author": "张恨水",
                                    "newRating": 950,
                                    "newRatingCount": 34,
                                },
                                "readingCount": 200,
                                "searchIdx": 9,
                            }
                        ],
                        "hasMore": False,
                    },
                }
            )

    warnings: list[str] = []
    cards = asyncio.run(
        weread_tools._fetch_gateway_publisher_book_cards(
            FakeClient(),
            "佳润互众",
            max_books=20,
            from_book_id="seed-book-id",
            warnings=warnings,
        )
    )

    assert warnings == []
    assert len(cards) == 2
    assert [call[1]["maxIdx"] for call in calls] == [0, 5]
    assert calls[0][0] == weread_tools.WEREAD_AGENT_GATEWAY
    assert calls[0][1]["api_name"] == "/store/search"
    assert calls[0][1]["type"] == 8
    assert calls[0][1]["skill_version"] == weread_tools.WEREAD_SKILL_VERSION
    assert calls[0][1]["fromBookId"] == "seed-book-id"
    assert calls[0][2]["Authorization"] == "Bearer wrk-test-key"
    assert cards[0]["source_kind"] == "agent_gateway_publisher"
    assert cards[0]["company_name"] == "佳润互众"


def test_fetch_app_publisher_book_cards_warns_when_cookie_is_rejected(monkeypatch) -> None:
    monkeypatch.delenv("WEREAD_API_KEY", raising=False)
    monkeypatch.setenv("WEREAD_COOKIE", "wr-test-cookie")

    class FakeResponse:
        status_code = 401

        def raise_for_status(self):
            return None

    class FakeClient:
        async def get(self, url, *, params=None, headers=None):
            return FakeResponse()

    warnings: list[str] = []
    cards = asyncio.run(
        weread_tools._fetch_app_publisher_book_cards(
            FakeClient(),
            "佳润互众",
            max_books=20,
            warnings=warnings,
        )
    )

    assert cards == []
    assert warnings
    assert "网页仍显示登录" in warnings[0]


def test_analyze_weread_company_books_filters_noise_by_detail_company(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(weread_tools, "WEREAD_COMPANY_INDEX_PATH", tmp_path / "weread_company_index.json")
    monkeypatch.setattr(weread_tools, "_COMPANY_INDEX_CACHE", None)

    async def fake_collect(client, company_name, *, max_books, **kwargs):
        return [
            {
                "title": "天涯神贴智慧大全",
                "author": "玄谷",
                "source_url": "https://weread.qq.com/web/reader/right",
                "company_name": None,
            },
            {
                "title": "佳偶天成",
                "author": "十四郎",
                "source_url": "https://weread.qq.com/web/reader/noise",
                "company_name": None,
            },
        ]

    async def fake_fetch_reader_result(client, item):
        if item.source_url.endswith("right"):
            return WeReadBookResult(
                query_title=item.title,
                query_author=item.author,
                matched_title=item.title,
                matched_author=item.author,
                recommendation=66.2,
                rating_count="222人点评",
                total_reading="4.1万人",
                reading_count="4.1万人",
                comment_count="222人点评",
                company_name="佳润互众",
                copyright_id="39149647",
                source_url=item.source_url,
                status="matched",
            )
        return WeReadBookResult(
            query_title=item.title,
            query_author=item.author,
            matched_title=item.title,
            matched_author=item.author,
            recommendation=82.9,
            company_name="别的公司",
            copyright_id="1000000190",
            source_url=item.source_url,
            status="matched",
        )

    monkeypatch.setattr(weread_tools, "_collect_company_book_candidates", fake_collect)
    monkeypatch.setattr(weread_tools, "fetch_reader_result", fake_fetch_reader_result)
    monkeypatch.setattr(weread_tools, "_attach_project_match", lambda result: None)

    rows = asyncio.run(weread_tools.analyze_weread_company_books("佳润互众"))

    assert len(rows) == 1
    assert rows[0].matched_title == "天涯神贴智慧大全"
    assert rows[0].company_name == "佳润互众"
    assert rows[0].copyright_id == "39149647"
    assert rows[0].reading_count == "4.1万人"
    assert rows[0].comment_count == "222人点评"


def test_analyze_weread_company_books_reuses_company_cache(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(weread_tools, "WEREAD_COMPANY_INDEX_PATH", tmp_path / "weread_company_index.json")
    monkeypatch.setattr(weread_tools, "_COMPANY_INDEX_CACHE", None)
    weread_tools._upsert_company_index_from_result(
        WeReadBookResult(
            query_title="天涯神贴智慧大全",
            query_author="玄谷",
            matched_title="天涯神贴智慧大全",
            matched_author="玄谷",
            recommendation=66.4,
            rating_count="223人点评",
            reading_count="4.1万人",
            company_name="佳润互众",
            copyright_id="39149647",
            source_url="https://weread.qq.com/web/reader/b3a32f00813aba926g018a42",
            status="matched",
        )
    )

    async def fake_fetch_reader_result(client, item):
        return WeReadBookResult(
            query_title=item.title,
            query_author=item.author,
            matched_title="天涯神贴智慧大全",
            matched_author="玄谷",
            recommendation=66.4,
            rating_count="223人点评",
            reading_count="4.1万人",
            company_name="佳润互众",
            copyright_id="39149647",
            source_url=item.source_url,
            status="matched",
        )

    async def fake_no_cards(*args, **kwargs):
        return []

    monkeypatch.setattr(weread_tools, "fetch_reader_result", fake_fetch_reader_result)
    monkeypatch.setattr(weread_tools, "_fetch_publisher_book_cards", fake_no_cards)
    monkeypatch.setattr(weread_tools, "_fetch_search_book_cards", fake_no_cards)
    monkeypatch.setattr(weread_tools, "_fetch_private_copyright_book_cards", fake_no_cards)
    monkeypatch.setattr(weread_tools, "_attach_project_match", lambda result: None)

    rows = asyncio.run(weread_tools.analyze_weread_company_books("佳润互众"))

    assert len(rows) == 1
    assert rows[0].matched_title == "天涯神贴智慧大全"
    assert rows[0].company_name == "佳润互众"
    assert rows[0].comment_count == "223人点评"


def test_analyze_weread_company_books_accepts_seed_title_copyright_id(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(weread_tools, "WEREAD_COMPANY_INDEX_PATH", tmp_path / "weread_company_index.json")
    monkeypatch.setattr(weread_tools, "_COMPANY_INDEX_CACHE", None)

    async def fake_search_cards(client, keyword, *, max_books):
        if keyword == "孟姜女（经典版）":
            return [
                {
                    "title": "孟姜女（经典版）",
                    "author": "张恨水",
                    "source_url": "https://weread.qq.com/web/bookDetail/3300187542",
                    "copyright_id": "39149647",
                    "recommendation": None,
                    "rating_count": "0人点评",
                    "reading_count": "1人",
                }
            ]
        return []

    async def fake_no_cards(*args, **kwargs):
        return []

    monkeypatch.setattr(weread_tools, "_fetch_search_book_cards", fake_search_cards)
    monkeypatch.setattr(weread_tools, "_fetch_publisher_book_cards", fake_no_cards)
    monkeypatch.setattr(weread_tools, "_fetch_private_copyright_book_cards", fake_no_cards)
    monkeypatch.setattr(weread_tools, "_attach_project_match", lambda result: None)

    rows = asyncio.run(weread_tools.analyze_weread_company_books("佳润互众\n孟姜女（经典版）"))

    assert len(rows) == 1
    assert rows[0].matched_title == "孟姜女（经典版）"
    assert rows[0].copyright_id == "39149647"
    assert rows[0].reading_count == "1人"


def test_extract_translator_from_weread_payload() -> None:
    body = '{"translator":"刘半农（译） 单华伟（编）","title":"1793乾隆英使觐见记"}'

    assert weread_tools._extract_translator(body) == "刘半农（译） 单华伟（编）"


def test_extract_nuxt_book_author_from_compressed_payload() -> None:
    body = 'window.__NUXT__=(function(a,b,c){return {data:[{bookInfo:{title:b,author:c}}]}}(0,"35岁以后，别急着认输","陈晨"));</script>'

    assert weread_tools._extract_nuxt_book_field(body, "title") == "35岁以后，别急着认输"
    assert weread_tools._extract_nuxt_book_field(body, "author") == "陈晨"


def test_extract_rating_count_normalizes_people_count() -> None:
    body = """
    <div class="new_book_rating_item_detail_count">
      <span><span class="rating_number_people">23</span>人点评</span>
    </div>
    """

    assert weread_tools._extract_rating_count(body) == "23人点评"
    assert weread_tools._normalize_comment_count("照微") is None
    assert weread_tools._normalize_comment_count("44") == "44人点评"


def test_project_url_key_uses_book_detail_id() -> None:
    url = "https://weread.qq.com/book-detail?type=1&senderVid=1&v=b5c32240813abb90ag01969c&wfrom=app"

    assert weread_tools._project_url_key(url) == "book-detail:b5c32240813abb90ag01969c"


def test_select_card_requires_author_match_when_author_is_provided() -> None:
    cards = [
        {"title": "同名书", "author": "错作者", "source_url": "https://weread.qq.com/web/reader/wrong"},
        {"title": "同名书", "author": "对作者", "source_url": "https://weread.qq.com/web/reader/right"},
    ]

    selected = weread_tools._select_card(cards, CatalogBook("同名书", "对作者"))

    assert selected is not None
    assert selected["source_url"] == "https://weread.qq.com/web/reader/right"
    assert weread_tools._select_card(cards[:1], CatalogBook("同名书", "对作者")) is None


def test_weread_export_contains_only_required_metric_columns() -> None:
    content = build_weread_xlsx(
        [
            {
                "query_title": "黄金投资指南",
                "query_author": "https://weread.qq.com/web/reader/demo123",
                "matched_title": "黄金投资指南",
                "translator": "刘半农（译）",
                "company_name": "广西师范大学出版社",
                "recommendation": 80.5,
                "reading_count": "623",
                "comment_count": "44",
                "project_editor": "单华伟",
                "project_title": "黄金投资指南",
                "project_source_sheet": "已完成书目",
                "project_match_status": "链接匹配",
                "status": "matched",
            }
        ]
    )
    shared_text = _xlsx_shared_text(content)

    assert "原始链接" in shared_text
    assert "原始作者" not in shared_text
    assert "https://weread.qq.com/web/reader/demo123" in shared_text
    assert "版权方（公司名字）" in shared_text
    assert "译者" in shared_text
    assert "阅读人数" in shared_text
    assert "评论数量" in shared_text
    assert "项目编辑" not in shared_text
    assert "项目书名" not in shared_text
    assert "利润" not in shared_text
    assert "来源链接" not in shared_text
    assert "匹配状态" not in shared_text
    assert "备注" not in shared_text
    assert "广西师范大学出版社" in shared_text
    assert "刘半农（译）" in shared_text


def test_weread_export_fills_empty_company_and_bad_comment_count() -> None:
    content = build_weread_xlsx(
        [
            {
                "query_title": "空版权方测试",
                "source_url": "https://weread.qq.com/book-detail?v=demo",
                "matched_title": "空版权方测试",
                "matched_author": "测试作者",
                "company_name": None,
                "recommendation": 0,
                "reading_count": "8人",
                "comment_count": "照微",
                "status": "matched",
            }
        ]
    )
    shared_text = _xlsx_shared_text(content)

    assert "无内容" in shared_text
    assert "0人点评" in shared_text
    assert ">照微<" not in shared_text


def test_analyze_weread_books_does_not_stop_at_80(monkeypatch) -> None:
    async def fake_fetch_search_result(client, item):
        return WeReadBookResult(
            query_title=item.title,
            query_author=item.author,
            matched_title=item.title,
            recommendation=80.0,
            status="matched",
        )

    monkeypatch.setattr(weread_tools, "fetch_search_result", fake_fetch_search_result)
    monkeypatch.setattr(weread_tools, "_attach_project_match", lambda result: None)

    catalog = "\n".join(f"测试书名{i}" for i in range(105))
    rows = asyncio.run(weread_tools.analyze_weread_books(catalog))

    assert len(rows) == 105


def test_weread_export_request_accepts_more_than_100_rows() -> None:
    row = WeReadBookRow(query_title="测试书名", status="matched")

    req = WeReadExportRequest(rows=[row for _ in range(120)])

    assert len(req.rows) == 120
